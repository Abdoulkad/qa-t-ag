#!/usr/bin/env python3
"""Analyse IA (LLM) + extraction SFD — moteur de l'app QA-T-AG.

L'analyse (revue, conception) est faite par un LLM (DeepSeek, API compatible
OpenAI). Les scripts déterministes contrôlent ensuite le résultat.
Principe : « le script calcule, l'agent analyse, l'humain arbitre » —
ici le LLM est la partie « agent », les skills sont ses prompts.
"""
import os

import docx

DEEPSEEK_URL = "https://api.deepseek.com/chat/completions"
MODEL = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
RACINE = os.path.dirname(os.path.abspath(__file__))


def extraire_sfd_texte(chemin):
    """Texte d'une SFD .docx ou .md (paragraphes + tableaux)."""
    if chemin.lower().endswith(".md"):
        with open(chemin, encoding="utf-8") as f:
            return f.read()
    d = docx.Document(chemin)
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n\n".join(parts)


def charger_prompt_skill(nom):
    """Corps du SKILL.md (sans frontmatter YAML) = prompt système de l'agent."""
    p = os.path.join(RACINE, "skills", nom, "SKILL.md")
    with open(p, encoding="utf-8") as f:
        lignes = f.read().splitlines()
    if lignes and lignes[0].strip() == "---":
        fin = next((i for i in range(1, len(lignes)) if lignes[i].strip() == "---"), 0)
        lignes = lignes[fin + 1:]
    return "\n".join(lignes)


def cle_disponible():
    return bool(os.environ.get("DEEPSEEK_API_KEY", "").strip())


def _log(msg):
    import datetime
    with open(os.path.join(RACINE, "ia.log"), "a", encoding="utf-8") as f:
        f.write(datetime.datetime.now().strftime("%H:%M:%S") + " " + msg + "\n")


def generer(system, user, max_tokens=8192):
    import time
    import requests
    k = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if not k:
        raise RuntimeError("DEEPSEEK_API_KEY non défini")
    _log("appel DeepSeek : modèle=%s, prompt=%d caractères"
         % (MODEL, len(system) + len(user)))
    t0 = time.time()
    r = requests.post(
        DEEPSEEK_URL,
        headers={"Authorization": "Bearer " + k},
        json={"model": MODEL,
              "messages": [{"role": "system", "content": system},
                           {"role": "user", "content": user}],
              "max_tokens": max_tokens},
        timeout=240,
    )
    _log("réponse en %.1fs, HTTP %s" % (time.time() - t0, r.status_code))
    if r.status_code != 200:
        _log("ERREUR : %s" % r.text[:300])
        raise RuntimeError("DeepSeek HTTP %s : %s" % (r.status_code, r.text[:300]))
    contenu = r.json()["choices"][0]["message"]["content"]
    _log("contenu reçu : %d caractères" % len(contenu))
    return contenu


def tester_cle():
    """Teste la clé DeepSeek (appel minimal). Retourne (ok, message)."""
    import requests
    k = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if not k:
        return False, "Aucune clé saisie"
    try:
        r = requests.post(DEEPSEEK_URL,
                          headers={"Authorization": "Bearer " + k},
                          json={"model": MODEL,
                                "messages": [{"role": "user", "content": "ping"}],
                                "max_tokens": 1},
                          timeout=30)
        if r.status_code == 200:
            return True, "Clé valide ✅ (modèle %s)" % MODEL
        return False, "HTTP %s — %s" % (r.status_code, r.text[:250])
    except Exception as e:
        return False, "Erreur réseau : %s" % e


def generer_revue(sfd_texte):
    return generer(charger_prompt_skill("qa-revue-sfd"),
                   "Voici la SFD à revoir :\n\n" + sfd_texte +
                   "\n\nContraintes de format STRICTES : "
                   "1) Relève au maximum 15 défauts, les plus importants "
                   "(bloquants et majeurs en priorité). "
                   "2) Inventaire REQ/BR concis (regroupe, une ligne par point). "
                   "3) Termine TOUJOURS par « ## Statut du livrable » "
                   "suivi d'une valeur canonique. "
                   "4) Chaque défaut a sa citation mot pour mot.")


def generer_cas(sfd_texte):
    return generer(charger_prompt_skill("qa-conception"),
                   "Voici la SFD validée :\n\n" + sfd_texte +
                   "\n\nContraintes STRICTES : sors UNIQUEMENT le CSV brut "
                   "(aucun texte avant ou après, aucun bloc de code, aucun commentaire). "
                   "La toute première ligne est l'en-tête exact : "
                   "ID;Titre;Objectif;Exigence;Technique;Priorité;Type;"
                   "Préconditions;Données de test;Étapes;Résultat attendu;Statut livrable. "
                   "Technique : partitions d'équivalence, valeurs limites, "
                   "table de décision ou transitions d'état (en minuscules). "
                   "Priorité : Critique, Haute ou Moyenne. "
                   "Type : Fonctionnel, Interface utilisateur ou Sécurité/Habilitation. "
                   "Statut livrable : OK ou À_CLARIFIER. "
                   "N'utilise JAMAIS le point-virgule « ; » à l'intérieur d'un champ "
                   "(dans Données de test ou Étapes, sépare par « / » ou « , »). "
                   "Numérote les CT séquentiellement (CT-001, CT-002, ...).")
