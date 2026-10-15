"""Génère le rapport de test (format DSI) à partir des données
d'exécution de QA-T-AG.

Usage :
    python generer_rapport.py TEMPLATE.docx rapport-input.json SORTIE.docx
                              [--interdits FICHIER]

Principe : le rapport est produit à partir des résultats, jamais rédigé
à la main. Le script refuse de générer (code de sortie 1) plutôt que de
produire un rapport incomplet ou non conforme :
  - statut ou priorité hors vocabulaire autorisé ;
  - priorité "Critique" sans validation humaine nommée ;
  - cas FAIL sans capture d'écran ;
  - cas FAIL, BLOQUÉ ou NON CONCLUANT sans commentaire ;
  - capture introuvable ;
  - valeur sensible (fichier --interdits) présente dans le rapport.

Les zones de signature (Rédacteur, Vérificateur, Approbateur) ne sont
jamais remplies : c'est un acte humain.

Dépendance : python-docx.
"""
import argparse
import copy
import json
import os
import sys

import docx
from docx.shared import Cm, RGBColor

NOIR = RGBColor(0, 0, 0)

# Vocabulaire de l'agent -> libellé du template
STATUTS = {
    "PASS": "PASS",
    "FAIL": "FAIL",
    "BLOCKED": "BLOQUÉ",
    "BLOQUÉ": "BLOQUÉ",
    "INCONCLUSIVE": "NON CONCLUANT",
    "NON CONCLUANT": "NON CONCLUANT",
    "NON EXÉCUTÉ": "NON EXÉCUTÉ",
}
RISQUE_VERS_PRIORITE = {"ÉLEVÉ": "Haute", "MOYEN": "Moyenne", "FAIBLE": "Basse"}
PRIORITES = {"Critique", "Haute", "Moyenne", "Basse"}
TYPES = {"Fonctionnel", "Sécurité / Habilitation", "Interface utilisateur",
         "Performance", "Accessibilité"}
STATUTS_AVEC_COMMENTAIRE = {"FAIL", "BLOQUÉ", "NON CONCLUANT"}


class ErreurRapport(Exception):
    pass


# --------------------------------------------------------------------------
# Validation des données d'entrée
# --------------------------------------------------------------------------
def valider(donnees, base_dir):
    erreurs = []
    for cas in donnees["cas"]:
        ident = cas.get("id", "?")
        statut = STATUTS.get(cas.get("statut", ""))
        if not statut:
            erreurs.append(f"{ident} : statut '{cas.get('statut')}' inconnu")
            continue
        cas["_statut"] = statut

        if "priorite" in cas:
            prio = cas["priorite"]
        else:
            prio = RISQUE_VERS_PRIORITE.get(cas.get("risque", ""))
        if prio not in PRIORITES:
            erreurs.append(f"{ident} : priorité/risque non reconnu")
        elif prio == "Critique" and not cas.get("critique_valide_par"):
            erreurs.append(f"{ident} : priorité Critique sans validation "
                           "humaine (champ critique_valide_par)")
        cas["_priorite"] = prio

        for t in cas.get("types", []):
            if t not in TYPES:
                erreurs.append(f"{ident} : type '{t}' inconnu")
        if not cas.get("types"):
            erreurs.append(f"{ident} : aucun type de test")

        captures = cas.get("captures", [])
        if statut == "FAIL" and not captures:
            erreurs.append(f"{ident} : FAIL sans capture d'écran")
        for c in captures:
            chemin = os.path.join(base_dir, c)
            if not os.path.isfile(chemin):
                erreurs.append(f"{ident} : capture introuvable ({c})")

        if statut in STATUTS_AVEC_COMMENTAIRE and not cas.get("commentaires"):
            erreurs.append(f"{ident} : statut {statut} sans commentaire")

        for champ in ("titre", "objectif", "exigence", "attendu", "observe",
                      "confiance", "testeur", "date"):
            if not cas.get(champ):
                erreurs.append(f"{ident} : champ '{champ}' vide")
    if erreurs:
        raise ErreurRapport("\n".join(erreurs))


# --------------------------------------------------------------------------
# Utilitaires de manipulation Word
# --------------------------------------------------------------------------
def supprimer(element):
    element.getparent().remove(element)


def ecrire_run(run, texte):
    run.text = texte
    run.font.italic = False
    run.font.color.rgb = NOIR


def cocher(paragraphe, libelles):
    """Remplace ☐ par ☒ devant chaque libellé sélectionné."""
    for run in paragraphe.runs:
        t = run.text
        if t.startswith("☐") and t[1:].replace("\u00a0", " ").strip() in libelles:
            run.text = "☒" + t[1:]


def vider_paragraphes_apres(cellule, garder):
    for p in cellule.paragraphs[garder:]:
        supprimer(p._p)


def paragraphe_apres(element, modele, texte, gras=False):
    """Insère un paragraphe (copie du format de `modele`) après `element`."""
    nouveau = copy.deepcopy(modele)
    for r in nouveau.findall(docx.oxml.ns.qn("w:r")):
        supprimer(r)
    element.addnext(nouveau)
    p = docx.text.paragraph.Paragraph(nouveau, None)
    run = p.add_run(texte)
    run.bold = gras
    return nouveau


# --------------------------------------------------------------------------
# Remplissage d'un bloc "cas de test"
# --------------------------------------------------------------------------
def remplir_cas(titre_p, table, cas, base_dir):
    # Titre du cas
    runs = titre_p.runs
    runs[0].text = f"{cas['id']} : {cas['titre']}"
    for r in runs[1:]:
        r.text = ""

    rows = table.rows

    # Objectif, données de test, étapes
    c = rows[1].cells[1]
    ps = c.paragraphs
    ecrire_run(ps[0].runs[1], cas["objectif"])
    ecrire_run(ps[1].runs[1], cas.get("donnees", "Aucune donnée saisie."))
    modele_etape = ps[4]._p
    supprimer(ps[3]._p)                       # consigne en italique
    for p in c.paragraphs[3:]:                # anciennes lignes 1. à 5.
        if p._p is not modele_etape:
            supprimer(p._p)
    precedent = modele_etape
    for i, etape in enumerate(cas.get("etapes", []), 1):
        nouveau = copy.deepcopy(modele_etape)
        precedent.addnext(nouveau)
        docx.text.paragraph.Paragraph(nouveau, None).runs[0].text = \
            f"{i}.  {etape}"
        precedent = nouveau
    supprimer(modele_etape)

    # Règle / exigence
    c = rows[2].cells[1]
    ecrire_run(c.paragraphs[0].runs[0], cas["exigence"])
    vider_paragraphes_apres(c, 1)

    # Priorité, type
    cocher(rows[3].cells[1].paragraphs[0], {cas["_priorite"]})
    cocher(rows[4].cells[1].paragraphs[0], set(cas["types"]))

    # Résultat attendu
    c = rows[5].cells[1]
    ecrire_run(c.paragraphs[0].runs[0], cas["attendu"])
    vider_paragraphes_apres(c, 1)

    # Résultat observé + confiance de preuve + référence d'exécution
    c = rows[6].cells[1]
    ecrire_run(c.paragraphs[0].runs[0], cas["observe"])
    vider_paragraphes_apres(c, 1)
    p = c.add_paragraph()
    r = p.add_run("Confiance de preuve : ")
    r.bold = True
    p.add_run(cas["confiance"])
    if cas.get("preuve"):
        p = c.add_paragraph()
        r = p.add_run("Preuve d'exécution : ")
        r.bold = True
        p.add_run(cas["preuve"])

    # Statut
    c = rows[7].cells[1]
    supprimer(c.paragraphs[0]._p)             # consigne "Cocher le statut..."
    cocher(c.paragraphs[0], {cas["_statut"]})

    # Captures
    c = rows[8].cells[1]
    vider_paragraphes_apres(c, 1)
    captures = cas.get("captures", [])
    if captures:
        for i, chemin in enumerate(captures):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph()
            p.add_run().add_picture(os.path.join(base_dir, chemin),
                                    width=Cm(12))
    else:
        c.paragraphs[0].add_run("Aucune capture (non requise pour ce statut).")

    # Testeur & date
    runs = rows[9].cells[1].paragraphs[0].runs
    runs[1].text = f"{cas['testeur']}          "
    runs[3].text = cas["date"]
    runs[4].text = ""
    runs[5].text = ""

    # Commentaires
    c = rows[10].cells[1]
    ecrire_run(c.paragraphs[0].runs[0], cas.get("commentaires") or "—")
    vider_paragraphes_apres(c, 1)


# --------------------------------------------------------------------------
# Génération
# --------------------------------------------------------------------------
def trouver_titre(doc, texte):
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == texte:
            return p
    raise ErreurRapport(f"Titre '{texte}' introuvable dans le template")


def generer(template, donnees, sortie, base_dir):
    valider(donnees, base_dir)
    doc = docx.Document(template)
    corps = doc.paragraphs
    modele_normal = next(p for p in corps
                         if p.text.startswith("Les objectifs"))._p

    # Organisation
    for p in corps:
        if "[NOM DE L" in p.text:
            p.runs[0].text = donnees["organisation"]
            for r in p.runs[1:]:
                r.text = ""
            break

    # Contexte
    ancre = trouver_titre(doc, "Contexte")._p
    for texte in donnees["contexte"]:
        ancre = paragraphe_apres(ancre, modele_normal, texte)

    # Objectifs
    ancre = next(p for p in doc.paragraphs
                 if p.text.startswith("Les objectifs"))._p
    for texte in donnees["objectifs"]:
        ancre = paragraphe_apres(ancre, modele_normal, "–  " + texte)

    # Exécution : retrait du guide et des consignes, introduction
    guide_tbl, cas_tbl = doc.tables[2], doc.tables[3]
    for p in list(doc.paragraphs):
        if (p.text.startswith("Guide d'utilisation")
                or p.text.startswith("Ce guide explique")
                or p.text.startswith("Modèle de cas de test")
                or p.text.startswith("Dupliquer ce bloc")):
            supprimer(p._p)
    precedent_guide = guide_tbl._tbl.getprevious()
    suivant_guide = guide_tbl._tbl.getnext()
    supprimer(guide_tbl._tbl)
    # paragraphe de liste vide qui suivait le guide (numéro orphelin)
    if suivant_guide is not None and not "".join(suivant_guide.itertext()).strip():
        supprimer(suivant_guide)

    ancre = trouver_titre(doc, "Exécution des tests")._p
    for texte in donnees.get("execution_intro", []):
        ancre = paragraphe_apres(ancre, modele_normal, texte)

    # Blocs de cas de test (duplication du modèle)
    titre_modele = cas_tbl._tbl.getprevious()
    table_modele = cas_tbl._tbl
    point = table_modele
    for cas in donnees["cas"]:
        titre = copy.deepcopy(titre_modele)
        table = copy.deepcopy(table_modele)
        point.addnext(titre)
        titre.addnext(table)
        point = table
        titre_p = docx.text.paragraph.Paragraph(titre, None)
        titre_p.paragraph_format.page_break_before = True
        remplir_cas(titre_p,
                    docx.table.Table(table, doc._body), cas, base_dir)
    supprimer(titre_modele)
    supprimer(table_modele)

    # Récapitulatif
    recap = doc.tables[-1]
    ligne_modele = recap.rows[1]._tr
    for i, ligne in enumerate(donnees["recapitulatif"], 1):
        tr = copy.deepcopy(ligne_modele)
        ligne_modele.addprevious(tr)
        cellules = docx.table._Row(tr, recap).cells
        valeurs = [str(i), ligne["test"], ligne["objectif"], ligne["attendu"],
                   ligne["obtenu"], ligne["impact"], ligne["observations"],
                   ligne["porteur"]]
        for cellule, valeur in zip(cellules, valeurs):
            cellule.paragraphs[0].text = valeur
    supprimer(ligne_modele)

    # Conclusion
    ancre = trouver_titre(doc, "Conclusion")._p
    for texte in donnees["conclusion"]:
        ancre = paragraphe_apres(ancre, modele_normal, texte)
    # paragraphes vides et saut de page final du template : page blanche
    suivant = ancre.getnext()
    while (suivant is not None
           and suivant.tag == docx.oxml.ns.qn("w:p")
           and not "".join(suivant.itertext()).strip()
           and suivant.find(".//" + docx.oxml.ns.qn("w:sectPr")) is None):
        a_supprimer, suivant = suivant, suivant.getnext()
        supprimer(a_supprimer)

    doc.core_properties.title = donnees.get("titre", "Rapport de test")
    doc.save(sortie)


def verifier_interdits(sortie, fichier):
    with open(fichier, encoding="utf-8") as f:
        interdits = [l.strip() for l in f if l.strip()]
    doc = docx.Document(sortie)
    textes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                textes.append(c.text)
    contenu = "\n".join(textes)
    trouves = [v for v in interdits if v in contenu]
    if trouves:
        os.remove(sortie)
        raise ErreurRapport(
            f"{len(trouves)} valeur(s) sensible(s) trouvée(s) dans le "
            "rapport : fichier supprimé, rien n'est diffusé.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("donnees")
    ap.add_argument("sortie")
    ap.add_argument("--interdits",
                    help="fichier listant les valeurs à ne jamais "
                         "écrire (une par ligne)")
    args = ap.parse_args()
    with open(args.donnees, encoding="utf-8") as f:
        donnees = json.load(f)
    base_dir = os.path.dirname(os.path.abspath(args.donnees))
    try:
        generer(args.template, donnees, args.sortie, base_dir)
        if args.interdits:
            verifier_interdits(args.sortie, args.interdits)
    except ErreurRapport as e:
        print("RAPPORT NON GÉNÉRÉ :\n" + str(e), file=sys.stderr)
        sys.exit(1)
    print("Rapport généré :", args.sortie)


if __name__ == "__main__":
    main()
