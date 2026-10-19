"""Extraction du texte d'une exigence déposée (.txt, .md, .docx, .pdf ; 10 Mo max)."""
import io

import docx
from pypdf import PdfReader

TAILLE_MAX = 10 * 1024 * 1024
EXTENSIONS = (".txt", ".md", ".docx", ".pdf")


class FichierRefuse(Exception):
    pass


def extraire(nom: str, contenu: bytes) -> str:
    if len(contenu) > TAILLE_MAX:
        raise FichierRefuse("Fichier trop volumineux (10 Mo maximum)")
    ext = "." + nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
    if ext not in EXTENSIONS:
        raise FichierRefuse("Format non pris en charge : %s (acceptés : %s)"
                            % (ext or "aucune extension", ", ".join(EXTENSIONS)))
    try:
        if ext in (".txt", ".md"):
            texte = contenu.decode("utf-8-sig")
        elif ext == ".docx":
            d = docx.Document(io.BytesIO(contenu))
            parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
                    parts.append(" | ".join(dict.fromkeys(cells)))
            texte = "\n\n".join(parts)
        else:
            lecteur = PdfReader(io.BytesIO(contenu))
            texte = "\n\n".join((p.extract_text() or "").strip() for p in lecteur.pages)
    except UnicodeDecodeError as e:
        raise FichierRefuse("Le fichier texte n'est pas encodé en UTF-8") from e
    except Exception as e:  # noqa: BLE001
        raise FichierRefuse("Fichier illisible : %s" % type(e).__name__) from e
    texte = texte.strip()
    if not texte:
        raise FichierRefuse("Aucun texte exploitable dans le fichier "
                            "(un PDF scanné doit d'abord passer par une reconnaissance de texte)")
    return texte
