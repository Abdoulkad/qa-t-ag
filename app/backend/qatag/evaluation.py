"""Outillage de la campagne d'évaluation des agents (protocole d'évaluation).

Le script calcule · l'agent analyse · l'humain arbitre : ce module exporte ce que
les agents ont réellement produit, propose un appariement avec le corrigé, puis
calcule les indicateurs I1 à I13 et la décision. Il ne décide jamais à la place
du relecteur : une ligne sans décision humaine bloque la consolidation.

    python -m qatag.evaluation modele-corrige corrige.xlsx
    python -m qatag.evaluation exporter campagne.xlsx --corrige corrige.xlsx [--reprendre ancienne.xlsx]
    python -m qatag.evaluation consolider campagne.xlsx --corrige corrige.xlsx synthese.xlsx

L'export lit la base désignée par DATABASE_URL, en lecture seule : la campagne se
joue sur une base dédiée, jamais sur la base de travail. Seules les missions dont
le titre suit EVAL-<exigence>-<config>-P<n> sont prises.
"""
import argparse
import random
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CATEGORIES = (
    "Contradiction certaine", "Contradiction apparente", "Ambiguïté",
    "Terme non défini", "Exigence non testable", "Règle manquante",
    "Cas limite à confirmer", "Règle sans exigence",
    "Incohérence avec les user stories",
)
CRITICITES = ("Bloquante", "Majeure", "Mineure")
DECISIONS = ("VP", "FP", "Doublon", "Trouvaille")
OUI_NON = ("oui", "non")
CAS_PAR_MISSION = 5
SEUIL_SIMILARITE = 0.6

RE_MISSION = re.compile(r"^EVAL-(?P<exigence>.+)-(?P<config>[^-\s]+)-P(?P<passage>\d+)\b")

# Indicateur : (libellé, opérateur, seuil, nature). Proposition de départ du
# protocole ; la valeur qui compte est celle de la feuille « Seuils » du corrigé,
# validée par la DSI et signée avant la campagne.
SEUILS = {
    "I1": ("Rappel (%)", ">=", 80, "Éliminatoire"),
    "I2": ("Bloquants oubliés", "<=", 0, "Éliminatoire"),
    "I8": ("Règles inventées acceptées", "<=", 0, "Éliminatoire"),
    "I13": ("Échecs techniques (%)", "<=", 5, "Éliminatoire"),
    "I3": ("Précision (%)", ">=", 70, "Qualité"),
    "I4": ("Criticité juste (%)", ">=", 75, "Qualité"),
    "I6": ("Citations valides (%)", ">=", 95, "Qualité"),
    "I9": ("Règles oubliées (%)", "<=", 10, "Qualité"),
    "I11": ("Cas justes (%)", ">=", 80, "Qualité"),
    "I12": ("Stabilité (écart moyen, points)", "<=", 15, "Qualité"),
}
INFORMATIFS = {
    "I5": "Catégorie juste (%)",
    "I7": "Trouvailles",
    "I10": "Couverture moyenne (%)",
}
ORDRE = ("I1", "I2", "I3", "I4", "I5", "I6", "I7", "I8", "I9", "I10", "I11", "I12", "I13")

# Colonnes des classeurs : relues par leur nom, l'ordre peut changer.
COL_MISSIONS = ("Mission", "Référence", "Exigence", "Config", "Passage", "Fournisseur", "Modèle",
                "Statut", "Défauts produits", "Règles produites", "Couverture %", "Cas produits",
                "Tâches lancées", "Tâches en échec", "Passage échoué", "Incident")
COL_DEFAUTS = ("Mission", "Exigence", "Config", "Passage", "Code", "Catégorie", "Criticité agent",
               "Citation", "Citation vérifiée", "Constat", "Proposition", "VT proposé", "Score",
               "Décision", "VT", "Commentaire")
COL_REGLES = ("Mission", "Exigence", "Config", "Passage", "Code", "Règle", "Statut appli",
              "Sources rejetées", "Règle du corrigé proposée", "Score", "Règle du corrigé",
              "Inventée", "Commentaire")
CRITERES_CAS = ("Vérifie le critère cité", "Données exécutables", "Résultat vérifiable",
                "Aucune règle ajoutée")
COL_CAS = ("Mission", "Exigence", "Config", "Passage", "Cas", "Titre", "Critère cité",
           "Données de test", "Résultat attendu") + CRITERES_CAS + ("Commentaire",)
SAISIE = {"Défauts": ("Décision", "VT", "Commentaire"),
          "Règles": ("Règle du corrigé", "Inventée", "Commentaire"),
          "Cas": CRITERES_CAS + ("Commentaire",),
          "Missions": ("Passage échoué", "Incident")}
CLES = {"Défauts": ("Mission", "Code"), "Règles": ("Mission", "Code"), "Cas": ("Mission", "Cas"),
        "Missions": ("Mission",)}


class ErreurEvaluation(Exception):
    """Classeur incomplet ou incohérent : la liste dit quoi corriger, ligne par ligne."""

    def __init__(self, erreurs):
        super().__init__("\n".join(erreurs))
        self.erreurs = list(erreurs)


# --- Noms de mission et similarité ------------------------------------------------

def lire_nom(titre: str):
    """EVAL-EX-01-DeepSeek-P2 → ("EX-01", "DeepSeek", 2) ; None hors campagne."""
    m = RE_MISSION.match((titre or "").strip())
    if not m:
        return None
    return m.group("exigence"), m.group("config"), int(m.group("passage"))


def _normaliser(texte: str) -> str:
    t = unicodedata.normalize("NFKD", texte or "")
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


MOTS_VIDES = set("""le la les un une des du de d l au aux et ou est sont etre doit doivent
peut peuvent son sa ses leur leurs ce cette ces qui que quoi dont par pour pendant avec sans
dans sur sous en il elle ils elles on ne pas plus tout tous toute toutes se si lors""".split())


def _racines(texte: str) -> set:
    """Mots porteurs de sens, réduits à leurs cinq premières lettres : « sécurisé »
    et « sécurité », « réinitialise » et « réinitialisation » se rejoignent."""
    return {m[:5] for m in _normaliser(texte).split() if len(m) > 2 and m not in MOTS_VIDES}


def similarite(a: str, b: str) -> float:
    """1 si un texte contient l'autre, sinon part des racines du plus court
    retrouvées dans l'autre (une règle reformulée reste reconnue)."""
    na, nb = _normaliser(a), _normaliser(b)
    if not na or not nb:
        return 0.0
    if na in nb or nb in na:
        return 1.0
    ra, rb = _racines(a), _racines(b)
    if not ra or not rb:
        return 0.0
    return round(len(ra & rb) / min(len(ra), len(rb)), 2)


def apparier(produits: list[dict], corrige: list[dict], champ: str, champ_corrige: str,
             cle_corrige: str) -> list[tuple]:
    """Propose, pour chaque élément produit, (proposition, id du corrigé, score).

    Le meilleur couple d'abord : chaque élément du corrigé reçoit au plus un VP ;
    un second élément qui vise le même passage est proposé en doublon ; le reste
    est proposé en FP. Ce n'est qu'une proposition : l'humain décide.
    """
    couples = sorted(((similarite(p[champ], c[champ_corrige]), i, c[cle_corrige])
                      for i, p in enumerate(produits) for c in corrige),
                     key=lambda x: (-x[0], x[1]))
    pris, propositions = set(), [None] * len(produits)
    for score, i, ident in couples:
        if score < SEUIL_SIMILARITE or propositions[i] is not None:
            continue
        if ident not in pris:
            pris.add(ident)
            propositions[i] = ("VP", ident, score)
    for score, i, ident in couples:
        if score >= SEUIL_SIMILARITE and propositions[i] is None:
            propositions[i] = ("Doublon", ident, score)
    return [p or ("FP", "", 0.0) for p in propositions]


# --- Lecture et écriture des classeurs ---------------------------------------------

ENTETE = PatternFill("solid", fgColor="1F5FA6")
A_SAISIR = PatternFill("solid", fgColor="FFF4CC")


def _feuille(wb, nom, colonnes, lignes, saisie=(), listes=None, largeurs=None):
    ws = wb.create_sheet(nom)
    ws.append(list(colonnes))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for ligne in lignes:
        ws.append([ligne.get(c, "") for c in colonnes])
    n = max(ws.max_row, 2)
    for j, col in enumerate(colonnes, start=1):
        lettre = ws.cell(row=1, column=j).column_letter
        ws.column_dimensions[lettre].width = (largeurs or {}).get(col, 14)
        if col in saisie:
            for i in range(2, n + 1):
                ws.cell(row=i, column=j).fill = A_SAISIR
        if listes and col in listes:
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(listes[col]),
                                allow_blank=True)
            dv.error = "Choisir une valeur de la liste."
            ws.add_data_validation(dv)
            dv.add("%s2:%s%d" % (lettre, lettre, max(n, 500)))
    ws.freeze_panes = "A2"
    return ws


def _lire(wb, nom, requis=True) -> list[dict]:
    if nom not in wb.sheetnames:
        if requis:
            raise ErreurEvaluation(["Feuille « %s » absente du classeur." % nom])
        return []
    lignes = wb[nom].iter_rows(values_only=True)
    entete = [str(c).strip() if c is not None else "" for c in next(lignes, [])]
    resultat = []
    for numero, valeurs in enumerate(lignes, start=2):
        if all(v is None or str(v).strip() == "" for v in valeurs):
            continue
        ligne = {k: ("" if v is None else (v.strip() if isinstance(v, str) else v))
                 for k, v in zip(entete, valeurs) if k}
        ligne["_ligne"] = numero
        resultat.append(ligne)
    return resultat


def _texte(v) -> str:
    return "" if v is None else str(v).strip()


def _lisez_moi(wb, lignes):
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.column_dimensions["A"].width = 110
    for i, t in enumerate(lignes, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if i == 1:
            c.font = Font(bold=True, size=14)


def modele_corrige(chemin) -> None:
    """Classeur vierge du corrigé : corpus, défauts VT, règles, seuils, signature."""
    wb = Workbook()
    _lisez_moi(wb, [
        "Corrigé de la campagne d'évaluation",
        "À remplir par les deux relecteurs QA, sans IA, avec la grille QA-REV-02, puis signer.",
        "Corpus : une ligne par exigence (EX-01…). Le même identifiant sert dans le titre des "
        "missions : EVAL-EX-01-DeepSeek-P1.",
        "Défauts : une ligne par défaut du corrigé (VT-01…), citation exacte du texte.",
        "Règles : les règles métier réellement présentes dans le texte (RC-01…).",
        "Seuils : valeurs validées par la DSI avant la campagne. Ne plus les modifier ensuite.",
        "Signature : le corrigé ne bouge plus une fois signé ; une trouvaille de l'agent est "
        "notée dans la campagne, jamais ajoutée ici.",
        "Ce classeur contient des exigences réelles : il ne va jamais dans le dépôt Git.",
    ])
    _feuille(wb, "Corpus", ("Exigence", "Type", "Titre", "Commentaire"), [],
             saisie=("Exigence", "Type", "Titre"),
             listes={"Type": ("User Story courte", "Extrait de SFD", "Calcul ou délai",
                              "Exigence claire")},
             largeurs={"Titre": 40, "Commentaire": 40, "Type": 20})
    _feuille(wb, "Défauts", ("Exigence", "VT", "Citation", "Catégorie", "Criticité", "Commentaire"),
             [], saisie=("Exigence", "VT", "Citation", "Catégorie", "Criticité"),
             listes={"Catégorie": CATEGORIES, "Criticité": CRITICITES},
             largeurs={"Citation": 50, "Catégorie": 26, "Commentaire": 40})
    _feuille(wb, "Règles", ("Exigence", "Règle", "Énoncé", "Commentaire"), [],
             saisie=("Exigence", "Règle", "Énoncé"), largeurs={"Énoncé": 70, "Commentaire": 30})
    _feuille(wb, "Seuils", ("Indicateur", "Libellé", "Opérateur", "Seuil", "Nature"),
             [{"Indicateur": k, "Libellé": v[0], "Opérateur": v[1], "Seuil": v[2], "Nature": v[3]}
              for k, v in SEUILS.items()], saisie=("Seuil",), largeurs={"Libellé": 34})
    _feuille(wb, "Signature", ("Champ", "Valeur"),
             [{"Champ": c} for c in ("Relecteur principal", "Second relecteur",
                                     "Date de signature du corrigé",
                                     "Date de validation des seuils (DSI)")],
             saisie=("Valeur",), largeurs={"Champ": 36, "Valeur": 30})
    wb.save(chemin)


def lire_corrige(chemin) -> dict:
    wb = load_workbook(chemin, data_only=True)
    corpus = _lire(wb, "Corpus")
    defauts = _lire(wb, "Défauts")
    regles = _lire(wb, "Règles")
    erreurs, avertissements = [], []
    exigences = {_texte(l.get("Exigence")) for l in corpus} - {""}
    if not exigences:
        erreurs.append("Corrigé, Corpus : aucune exigence.")

    def verifier(lignes, feuille, cle, valeurs=None):
        vus = set()
        for l in lignes:
            ex, ident = _texte(l.get("Exigence")), _texte(l.get(cle))
            ou = "Corrigé, %s ligne %d" % (feuille, l["_ligne"])
            if ex not in exigences:
                erreurs.append("%s : exigence « %s » absente du Corpus." % (ou, ex))
            if not ident:
                erreurs.append("%s : %s manquant." % (ou, cle))
            elif (ex, ident) in vus:
                erreurs.append("%s : %s en double pour %s." % (ou, ident, ex))
            vus.add((ex, ident))
            for col, permises in (valeurs or {}).items():
                if _texte(l.get(col)) not in permises:
                    erreurs.append("%s : %s « %s » hors liste." % (ou, col, _texte(l.get(col))))

    verifier(defauts, "Défauts", "VT", {"Catégorie": CATEGORIES, "Criticité": CRITICITES})
    verifier(regles, "Règles", "Règle")
    seuils = {}
    for l in _lire(wb, "Seuils"):
        k = _texte(l.get("Indicateur"))
        if k not in SEUILS:
            continue
        try:
            seuils[k] = (SEUILS[k][0], SEUILS[k][1], float(l.get("Seuil")), SEUILS[k][3])
        except (TypeError, ValueError):
            erreurs.append("Corrigé, Seuils ligne %d : seuil de %s non numérique." % (l["_ligne"], k))
    for k in SEUILS:
        if k not in seuils and not any(k in e for e in erreurs):
            erreurs.append("Corrigé, Seuils : %s manquant." % k)
    signature = {_texte(l.get("Champ")): _texte(l.get("Valeur")) for l in _lire(wb, "Signature")}
    for champ, valeur in signature.items():
        if not valeur:
            avertissements.append("Corrigé non signé : « %s » est vide." % champ)
    if erreurs:
        raise ErreurEvaluation(erreurs)
    return {
        "exigences": sorted(exigences),
        "defauts": [{"exigence": _texte(l["Exigence"]), "vt": _texte(l["VT"]),
                     "citation": _texte(l.get("Citation")), "categorie": _texte(l["Catégorie"]),
                     "criticite": _texte(l["Criticité"])} for l in defauts],
        "regles": [{"exigence": _texte(l["Exigence"]), "regle": _texte(l["Règle"]),
                    "enonce": _texte(l.get("Énoncé"))} for l in regles],
        "seuils": seuils, "signature": signature, "avertissements": avertissements,
    }


# --- Extraction depuis la base --------------------------------------------------------

def _passage_echoue(taches) -> bool:
    """Relancé une fois puis réussi : simple incident. Deux échecs pour une même
    étape, ou une étape restée en échec : passage échoué (résultat à part entière)."""
    par_type = defaultdict(list)
    for t in sorted(taches, key=lambda t: t.id):
        par_type[t.type].append(t.etat)
    return any(e.count("ECHEC") >= 2 or e[-1] == "ECHEC" for e in par_type.values())


def extraire(db, prefixe: str = "EVAL-") -> list[dict]:
    """Sorties brutes des agents pour chaque mission de la campagne, lues en base."""
    from sqlalchemy import select

    from . import missions as services
    from .models import CasTest, Mission, TacheIA

    resultat = []
    for m in db.scalars(select(Mission).where(Mission.titre.startswith(prefixe))
                        .order_by(Mission.titre)):
        nom = lire_nom(m.titre)
        if not nom or not m.versions:
            continue
        v = m.versions[-1]
        taches = list(db.scalars(select(TacheIA).where(TacheIA.mission_id == m.id)))
        phrases = [p for p in services.couverture(v) if p["statut"] not in ("TITRE", "CADRE")]
        couverture = (round(100 * sum(p["statut"] == "COUVERTE" for p in phrases) / len(phrases))
                      if phrases and v.regles else None)
        cas = list(db.scalars(select(CasTest).where(CasTest.version_id == v.id)
                              .order_by(CasTest.code)))
        resultat.append({
            "mission": m.titre, "reference": m.reference, "exigence": nom[0], "config": nom[1],
            "passage": nom[2], "fournisseur": v.fournisseur_ia or "", "modele": v.modele_ia or "",
            "statut": m.statut, "couverture": couverture,
            "taches": len(taches), "echecs": sum(t.etat == "ECHEC" for t in taches),
            "echoue": _passage_echoue(taches),
            "defauts": [{"code": d.code, "categorie": d.categorie,
                         "criticite": d.gravite_initiale or d.gravite, "citation": d.citation or "",
                         "citation_verifiee": bool(d.citation_verifiee), "constat": d.constat or ""}
                        for d in v.defauts],
            "regles": [{"code": r.code, "texte": r.texte, "statut": r.statut,
                        "rejetees": "; ".join(r.sources_rejetees or [])} for r in v.regles],
            "cas": [{"code": c.code, "titre": c.titre, "critere": c.exigence,
                     "donnees": c.donnees_test, "attendu": c.resultat_attendu} for c in cas],
        })
    return resultat


def echantillon(mission: str, cas: list[dict], n: int = CAS_PAR_MISSION) -> list[dict]:
    """Cas tirés au hasard, mais toujours les mêmes pour une mission donnée :
    un nouvel export ne change pas l'échantillon déjà jugé."""
    if len(cas) <= n:
        return list(cas)
    choisis = random.Random(mission).sample(range(len(cas)), n)
    return [cas[i] for i in sorted(choisis)]


# --- Classeur de campagne ---------------------------------------------------------------

def _reprises(chemin) -> dict:
    """Décisions déjà saisies dans un export précédent, par feuille et par clé."""
    if not chemin:
        return {}
    wb = load_workbook(chemin, data_only=True)
    reprises = {}
    for feuille, colonnes in SAISIE.items():
        for l in _lire(wb, feuille, requis=False):
            cle = tuple(_texte(l.get(c)) for c in CLES[feuille])
            valeurs = {c: l.get(c, "") for c in colonnes if _texte(l.get(c))}
            if valeurs:
                reprises[(feuille, cle)] = valeurs
    return reprises


def ecrire_campagne(missions: list[dict], corrige: dict, chemin, reprendre=None) -> dict:
    """Classeur de saisie : faits extraits, propositions d'appariement, colonnes
    jaunes à remplir. Renvoie quelques comptes pour l'affichage."""
    reprises = _reprises(reprendre)
    inconnues = sorted({m["exigence"] for m in missions} - set(corrige["exigences"]))
    if inconnues:
        raise ErreurEvaluation(["Missions dont l'exigence n'est pas dans le Corpus du corrigé : "
                                + ", ".join(inconnues)])

    def avec_reprise(feuille, ligne):
        cle = tuple(_texte(ligne.get(c)) for c in CLES[feuille])
        ligne.update(reprises.get((feuille, cle), {}))
        return ligne

    lignes_m, lignes_d, lignes_r, lignes_c = [], [], [], []
    for m in missions:
        base = {"Mission": m["mission"], "Exigence": m["exigence"], "Config": m["config"],
                "Passage": "P%d" % m["passage"]}
        lignes_m.append(avec_reprise("Missions", {
            **base, "Référence": m["reference"], "Fournisseur": m["fournisseur"],
            "Modèle": m["modele"], "Statut": m["statut"], "Défauts produits": len(m["defauts"]),
            "Règles produites": len(m["regles"]),
            "Couverture %": "" if m["couverture"] is None else m["couverture"],
            "Cas produits": len(m["cas"]), "Tâches lancées": m["taches"],
            "Tâches en échec": m["echecs"], "Passage échoué": "oui" if m["echoue"] else "non"}))
        vt = [d for d in corrige["defauts"] if d["exigence"] == m["exigence"]]
        for d, (prop, ident, score) in zip(m["defauts"], apparier(
                m["defauts"], vt, "citation", "citation", "vt")):
            lignes_d.append(avec_reprise("Défauts", {
                **base, "Code": d["code"], "Catégorie": d["categorie"],
                "Criticité agent": d["criticite"], "Citation": d["citation"],
                "Citation vérifiée": "oui" if d["citation_verifiee"] else "non",
                "Constat": d["constat"], "Proposition": prop, "VT proposé": ident,
                "Score": score}))
        rc = [r for r in corrige["regles"] if r["exigence"] == m["exigence"]]
        for r, (prop, ident, score) in zip(m["regles"], apparier(
                m["regles"], rc, "texte", "enonce", "regle")):
            lignes_r.append(avec_reprise("Règles", {
                **base, "Code": r["code"], "Règle": r["texte"], "Statut appli": r["statut"],
                "Sources rejetées": r["rejetees"],
                "Règle du corrigé proposée": ident if prop == "VP" else "",
                "Score": score if prop == "VP" else ""}))
        for c in echantillon(m["mission"], m["cas"]):
            lignes_c.append(avec_reprise("Cas", {
                **base, "Cas": c["code"], "Titre": c["titre"], "Critère cité": c["critere"],
                "Données de test": c["donnees"], "Résultat attendu": c["attendu"]}))

    wb = Workbook()
    _lisez_moi(wb, [
        "Campagne d'évaluation — grille de saisie",
        "Exporté le %s depuis la base d'évaluation. Colonnes jaunes : à remplir par le "
        "relecteur. Le reste vient de l'application, sans retouche." % date.today().isoformat(),
        "Défauts : la Proposition et le VT proposé viennent d'un rapprochement des citations. "
        "Ce n'est qu'une aide : la Décision (VP, FP, Doublon, Trouvaille) et le VT sont les "
        "tiens. Un VP ou un Doublon cite son VT ; une Trouvaille exige l'accord des deux "
        "relecteurs, sinon c'est un FP.",
        "Les oublis (FN) ne se saisissent pas : tout défaut du corrigé qu'aucun VP ne cite "
        "dans une mission est compté oublié pour cette mission.",
        "Règles : renseigner la règle du corrigé reprise, ou à défaut Inventée oui/non "
        "(non si la règle vient d'une réponse métier écrite d'avance).",
        "Cas : quatre conditions oui/non ; un cas est juste si les quatre sont à oui.",
        "Missions : Passage échoué est calculé (deux échecs d'une même étape, ou une étape "
        "restée en échec) ; corriger seulement si l'incident le justifie, avec un motif.",
        "Un nouvel export avec --reprendre conserve tout ce qui a déjà été saisi.",
        "Ce classeur contient des exigences réelles : il ne va jamais dans le dépôt Git.",
    ])
    _feuille(wb, "Missions", COL_MISSIONS, lignes_m, saisie=SAISIE["Missions"],
             listes={"Passage échoué": OUI_NON},
             largeurs={"Mission": 30, "Modèle": 22, "Incident": 40})
    _feuille(wb, "Défauts", COL_DEFAUTS, lignes_d, saisie=SAISIE["Défauts"],
             listes={"Décision": DECISIONS},
             largeurs={"Mission": 30, "Citation": 45, "Constat": 45, "Catégorie": 24,
                       "Commentaire": 30})
    _feuille(wb, "Règles", COL_REGLES, lignes_r, saisie=SAISIE["Règles"],
             listes={"Inventée": OUI_NON},
             largeurs={"Mission": 30, "Règle": 60, "Sources rejetées": 30, "Commentaire": 30})
    _feuille(wb, "Cas", COL_CAS, lignes_c, saisie=SAISIE["Cas"],
             listes={c: OUI_NON for c in CRITERES_CAS},
             largeurs={"Mission": 30, "Titre": 36, "Critère cité": 36, "Données de test": 36,
                       "Résultat attendu": 36, "Commentaire": 30})
    wb.save(chemin)
    return {"missions": len(lignes_m), "defauts": len(lignes_d), "regles": len(lignes_r),
            "cas": len(lignes_c), "reprises": len(reprises)}


# --- Consolidation -----------------------------------------------------------------------

def _pct(n, d):
    return None if not d else round(100 * n / d, 1)


def _passe(valeur, operateur, seuil):
    if valeur is None:
        return False
    return valeur >= seuil if operateur == ">=" else valeur <= seuil


def lire_campagne(chemin) -> dict:
    wb = load_workbook(chemin, data_only=True)
    return {f: _lire(wb, f) for f in ("Missions", "Défauts", "Règles", "Cas")}


def consolider(campagne: dict, corrige: dict) -> dict:
    """Indicateurs par configuration. Refuse de calculer tant qu'une décision manque."""
    erreurs = []
    vt_par_ex = defaultdict(dict)
    for d in corrige["defauts"]:
        vt_par_ex[d["exigence"]][d["vt"]] = d
    rc_par_ex = defaultdict(set)
    for r in corrige["regles"]:
        rc_par_ex[r["exigence"]].add(r["regle"])

    missions = {}
    for l in campagne["Missions"]:
        nom = _texte(l.get("Mission"))
        echoue = _texte(l.get("Passage échoué")).lower()
        if echoue not in OUI_NON:
            erreurs.append("Missions ligne %d : Passage échoué doit valoir oui ou non." % l["_ligne"])
        if _texte(l.get("Exigence")) not in corrige["exigences"]:
            erreurs.append("Missions ligne %d : exigence « %s » absente du corrigé."
                           % (l["_ligne"], _texte(l.get("Exigence"))))
        couv = l.get("Couverture %")
        missions[nom] = {"exigence": _texte(l.get("Exigence")), "config": _texte(l.get("Config")),
                         "passage": _texte(l.get("Passage")), "echoue": echoue == "oui",
                         "couverture": couv if isinstance(couv, (int, float)) else None,
                         "vp": set(), "fp": 0, "doublons": 0, "trouvailles": [],
                         "crit_ok": 0, "cat_ok": 0, "cit": 0, "cit_ok": 0, "confusions": [],
                         "rc": set(), "inventees": 0, "inventees_rejetees": 0,
                         "cas": 0, "cas_justes": 0}

    def mission_de(l, feuille):
        m = missions.get(_texte(l.get("Mission")))
        if m is None:
            erreurs.append("%s ligne %d : mission « %s » absente de la feuille Missions."
                           % (feuille, l["_ligne"], _texte(l.get("Mission"))))
        return m

    for l in campagne["Défauts"]:
        m = mission_de(l, "Défauts")
        if m is None:
            continue
        ou = "Défauts ligne %d (%s %s)" % (l["_ligne"], _texte(l.get("Mission")),
                                           _texte(l.get("Code")))
        if _texte(l.get("Citation")):
            m["cit"] += 1
            m["cit_ok"] += _texte(l.get("Citation vérifiée")).lower() == "oui"
        decision, vt = _texte(l.get("Décision")), _texte(l.get("VT"))
        if decision not in DECISIONS:
            erreurs.append("%s : Décision manquante ou hors liste (VP, FP, Doublon, Trouvaille)." % ou)
            continue
        if decision in ("VP", "Doublon"):
            ref = vt_par_ex[m["exigence"]].get(vt)
            if ref is None:
                erreurs.append("%s : %s sans VT valable pour %s (« %s »)."
                               % (ou, decision, m["exigence"], vt))
                continue
            if decision == "Doublon":
                m["doublons"] += 1
                continue
            if vt in m["vp"]:
                erreurs.append("%s : %s déjà crédité d'un VP dans cette mission ; le second est "
                               "un Doublon." % (ou, vt))
                continue
            m["vp"].add(vt)
            m["crit_ok"] += _texte(l.get("Criticité agent")) == ref["criticite"]
            if _texte(l.get("Catégorie")) == ref["categorie"]:
                m["cat_ok"] += 1
            else:
                m["confusions"].append((ref["categorie"], _texte(l.get("Catégorie"))))
        elif decision == "FP":
            m["fp"] += 1
        else:
            m["trouvailles"].append({"citation": _texte(l.get("Citation")),
                                     "categorie": _texte(l.get("Catégorie")),
                                     "constat": _texte(l.get("Constat")),
                                     "commentaire": _texte(l.get("Commentaire"))})

    for l in campagne["Règles"]:
        m = mission_de(l, "Règles")
        if m is None:
            continue
        ou = "Règles ligne %d (%s %s)" % (l["_ligne"], _texte(l.get("Mission")),
                                          _texte(l.get("Code")))
        rc, inventee = _texte(l.get("Règle du corrigé")), _texte(l.get("Inventée")).lower()
        if rc:
            if rc not in rc_par_ex[m["exigence"]]:
                erreurs.append("%s : règle du corrigé « %s » inconnue pour %s."
                               % (ou, rc, m["exigence"]))
            m["rc"].add(rc)
        elif inventee not in OUI_NON:
            erreurs.append("%s : renseigner la Règle du corrigé ou Inventée (oui/non)." % ou)
        elif inventee == "oui":
            if _texte(l.get("Statut appli")) == "HYPOTHESE":
                m["inventees_rejetees"] += 1
            else:
                m["inventees"] += 1

    for l in campagne["Cas"]:
        m = mission_de(l, "Cas")
        if m is None:
            continue
        valeurs = [_texte(l.get(c)).lower() for c in CRITERES_CAS]
        if any(v not in OUI_NON for v in valeurs):
            erreurs.append("Cas ligne %d (%s %s) : les quatre conditions doivent valoir oui ou non."
                           % (l["_ligne"], _texte(l.get("Mission")), _texte(l.get("Cas"))))
            continue
        m["cas"] += 1
        m["cas_justes"] += all(v == "oui" for v in valeurs)

    if erreurs:
        raise ErreurEvaluation(erreurs)

    configs = defaultdict(list)
    for nom, m in missions.items():
        configs[m["config"]].append((nom, m))

    resultats = {}
    for config, liste in sorted(configs.items()):
        c = Counter()
        bloquants, trouvailles, confusions = [], [], Counter()
        rappel_par_ex = defaultdict(list)
        couvertures = []
        for nom, m in liste:
            vt = vt_par_ex[m["exigence"]]
            fn = [v for k, v in vt.items() if k not in m["vp"]]
            c["vp"] += len(m["vp"])
            c["fn"] += len(fn)
            c["fp"] += m["fp"]
            c["doublons"] += m["doublons"]
            c["trouvailles"] += len(m["trouvailles"])
            c["crit_ok"] += m["crit_ok"]
            c["cat_ok"] += m["cat_ok"]
            c["cit"] += m["cit"]
            c["cit_ok"] += m["cit_ok"]
            c["rc"] += len(rc_par_ex[m["exigence"]])
            c["rc_oubliees"] += len(rc_par_ex[m["exigence"]] - m["rc"])
            c["inventees"] += m["inventees"]
            c["inventees_rejetees"] += m["inventees_rejetees"]
            c["cas"] += m["cas"]
            c["cas_justes"] += m["cas_justes"]
            c["passages"] += 1
            c["echecs"] += m["echoue"]
            confusions.update(m["confusions"])
            if m["couverture"] is not None:
                couvertures.append(m["couverture"])
            if vt:
                rappel_par_ex[m["exigence"]].append(100 * len(m["vp"]) / len(vt))
            for v in fn:
                if v["criticite"] == "Bloquante":
                    bloquants.append({"mission": nom, "passage": m["passage"], **v})
            trouvailles += [{"mission": nom, "exigence": m["exigence"], **t}
                            for t in m["trouvailles"]]
        ecarts = {ex: round(max(r) - min(r), 1) for ex, r in rappel_par_ex.items() if len(r) >= 2}
        ind = {
            "I1": _pct(c["vp"], c["vp"] + c["fn"]),
            "I2": len(bloquants),
            "I3": _pct(c["vp"], c["vp"] + c["fp"]),
            "I4": _pct(c["crit_ok"], c["vp"]),
            "I5": _pct(c["cat_ok"], c["vp"]),
            "I6": _pct(c["cit_ok"], c["cit"]),
            "I7": c["trouvailles"],
            "I8": c["inventees"],
            "I9": _pct(c["rc_oubliees"], c["rc"]) if c["rc"] else 0.0,
            "I10": round(sum(couvertures) / len(couvertures), 1) if couvertures else None,
            "I11": _pct(c["cas_justes"], c["cas"]),
            "I12": round(sum(ecarts.values()) / len(ecarts), 1) if ecarts else None,
            "I13": _pct(c["echecs"], c["passages"]),
        }
        verdicts = {k: _passe(ind[k], s[1], s[2]) for k, s in corrige["seuils"].items()}
        elim = all(verdicts[k] for k, s in corrige["seuils"].items() if s[3] == "Éliminatoire")
        qualite = sum(verdicts[k] for k, s in corrige["seuils"].items() if s[3] == "Qualité")
        decision = ("Non apte" if not elim or qualite < 2
                    else "Apte" if qualite >= 4 else "Apte sous réserve")
        resultats[config] = {
            "comptes": dict(c), "indicateurs": ind, "verdicts": verdicts,
            "qualite": qualite, "eliminatoires": elim, "decision": decision,
            "bloquants": bloquants, "trouvailles": trouvailles,
            "confusions": confusions.most_common(), "ecarts": ecarts,
            "modeles": sorted({_texte(l.get("Modèle")) for l in campagne["Missions"]
                               if _texte(l.get("Config")) == config} - {""}),
        }
    avertissements = list(corrige["avertissements"])
    for config, r in resultats.items():
        if len(r["modeles"]) > 1:   # étape 1 du protocole : configuration gelée
            avertissements.append("Configuration %s : plusieurs modèles dans la campagne (%s)."
                                  % (config, ", ".join(r["modeles"])))
    return {"configs": resultats, "avertissements": avertissements}


def _nombre(v) -> str:
    """66.7 → « 66,7 » ; 80.0 → « 80 » : écriture française, sans décimale inutile."""
    v = round(float(v), 1)
    return ("%d" % v if v == int(v) else ("%.1f" % v).replace(".", ","))


def _format(k, v) -> str:
    if v is None:
        return "non mesurable"
    if k in ("I2", "I7", "I8"):
        return "%d" % v
    if k == "I12":
        return "%s pts" % _nombre(v)
    return "%s %%" % _nombre(v)


def ecrire_synthese(resultat: dict, corrige: dict, chemin) -> None:
    configs = resultat["configs"]
    noms = list(configs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws.append(["Indicateur", *noms, "Seuil", "Nature"])
    for k in ORDRE:
        if k in corrige["seuils"]:
            lib, op, seuil, nature = corrige["seuils"][k]
            cellules = ["%s %s" % (_format(k, configs[n]["indicateurs"][k]),
                                   "✓" if configs[n]["verdicts"][k] else "✗") for n in noms]
            ws.append(["%s %s" % (k, lib), *cellules, "%s %s" % ("≥" if op == ">=" else "≤",
                                                                 _format(k, seuil)), nature])
        else:
            ws.append(["%s %s" % (k, INFORMATIFS[k]),
                       *[_format(k, configs[n]["indicateurs"][k]) for n in noms], "", "Information"])
    ws.append(["Seuils de qualité atteints", *["%d / 6" % configs[n]["qualite"] for n in noms]])
    ws.append(["Décision", *[configs[n]["decision"] for n in noms]])
    ws.append(["Modèle", *[", ".join(configs[n]["modeles"]) for n in noms]])
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(horizontal="left")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
    for c in ws[ws.max_row - 1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    for j in range(2, len(noms) + 4):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 20

    libelles = {"vp": "VP", "fp": "FP", "fn": "FN (oublis)", "doublons": "Doublons",
                "trouvailles": "Trouvailles", "crit_ok": "VP de criticité juste",
                "cat_ok": "VP de catégorie juste", "cit": "Citations produites",
                "cit_ok": "Citations vérifiées", "rc": "Règles du corrigé attendues",
                "rc_oubliees": "Règles du corrigé oubliées",
                "inventees": "Règles inventées acceptées",
                "inventees_rejetees": "Règles inventées rejetées (hypothèse)",
                "cas": "Cas jugés", "cas_justes": "Cas justes", "passages": "Passages",
                "echecs": "Passages échoués"}
    _feuille(wb, "Comptages", ("Comptage", *noms),
             [{"Comptage": lib, **{n: configs[n]["comptes"].get(k, 0) for n in noms}}
              for k, lib in libelles.items()], largeurs={"Comptage": 38})
    _feuille(wb, "Bloquants oubliés", ("Config", "Mission", "Passage", "VT", "Citation",
                                       "Catégorie"),
             [{"Config": n, "Mission": b["mission"], "Passage": b["passage"], "VT": b["vt"],
               "Citation": b["citation"], "Catégorie": b["categorie"]}
              for n in noms for b in configs[n]["bloquants"]],
             largeurs={"Mission": 30, "Citation": 60, "Catégorie": 24})
    _feuille(wb, "Trouvailles", ("Config", "Mission", "Citation", "Catégorie", "Constat",
                                 "Commentaire"),
             [{"Config": n, "Mission": t["mission"], "Citation": t["citation"],
               "Catégorie": t["categorie"], "Constat": t["constat"],
               "Commentaire": t["commentaire"]} for n in noms for t in configs[n]["trouvailles"]],
             largeurs={"Mission": 30, "Citation": 45, "Constat": 50, "Commentaire": 30})
    _feuille(wb, "Catégories confondues", ("Config", "Catégorie du corrigé", "Catégorie de l'agent",
                                           "Nombre"),
             [{"Config": n, "Catégorie du corrigé": a, "Catégorie de l'agent": b, "Nombre": k}
              for n in noms for (a, b), k in configs[n]["confusions"]],
             largeurs={"Catégorie du corrigé": 30, "Catégorie de l'agent": 30})
    _feuille(wb, "Stabilité", ("Config", "Exigence", "Écart de rappel (points)"),
             [{"Config": n, "Exigence": ex, "Écart de rappel (points)": e}
              for n in noms for ex, e in sorted(configs[n]["ecarts"].items())],
             largeurs={"Écart de rappel (points)": 26})
    _feuille(wb, "Avertissements", ("Avertissement",),
             [{"Avertissement": a} for a in resultat["avertissements"]] or
             [{"Avertissement": "Aucun"}], largeurs={"Avertissement": 90})
    wb.save(chemin)


# --- Ligne de commande ---------------------------------------------------------------------

def _afficher(resultat):
    for nom, r in resultat["configs"].items():
        print("\n%s — %s" % (nom, r["decision"]))
        for k in ORDRE:
            v = _format(k, r["indicateurs"][k])
            marque = ("" if k not in r["verdicts"] else " ✓" if r["verdicts"][k] else " ✗")
            print("  %-4s %-36s %s%s" % (k, SEUILS.get(k, (INFORMATIFS.get(k),))[0], v, marque))
    for a in resultat["avertissements"]:
        print("Attention : " + a)


def principal(argv=None) -> int:
    p = argparse.ArgumentParser(prog="python -m qatag.evaluation",
                                description="Campagne d'évaluation des agents QA-T-AG.")
    sous = p.add_subparsers(dest="commande", required=True)
    s = sous.add_parser("modele-corrige", help="crée le classeur vierge du corrigé")
    s.add_argument("sortie")
    s = sous.add_parser("exporter", help="exporte les missions EVAL-* de la base")
    s.add_argument("sortie")
    s.add_argument("--corrige", required=True)
    s.add_argument("--reprendre", help="export précédent dont on garde les saisies")
    s.add_argument("--prefixe", default="EVAL-")
    s = sous.add_parser("consolider", help="calcule les indicateurs et la décision")
    s.add_argument("campagne")
    s.add_argument("sortie")
    s.add_argument("--corrige", required=True)
    a = p.parse_args(argv)

    try:
        if a.commande == "modele-corrige":
            if Path(a.sortie).exists():
                print("%s existe déjà : choisir un autre nom." % a.sortie)
                return 1
            modele_corrige(a.sortie)
            print("Corrigé vierge créé : %s" % a.sortie)
        elif a.commande == "exporter":
            if Path(a.sortie).exists():
                print("%s existe déjà : choisir un autre nom (les saisies se reprennent "
                      "avec --reprendre)." % a.sortie)
                return 1
            from sqlalchemy.engine import make_url

            from .config import parametres
            from .db import SessionLocale
            corrige = lire_corrige(a.corrige)
            print("Base lue : %s" % make_url(parametres().database_url).database)
            with SessionLocale() as db:
                missions = extraire(db, a.prefixe)
            if not missions:
                print("Aucune mission dont le titre commence par %s." % a.prefixe)
                return 1
            n = ecrire_campagne(missions, corrige, a.sortie, a.reprendre)
            print("Campagne exportée : %s — %d missions, %d défauts, %d règles, %d cas à juger%s."
                  % (a.sortie, n["missions"], n["defauts"], n["regles"], n["cas"],
                     ", %d saisies reprises" % n["reprises"] if a.reprendre else ""))
        else:
            if Path(a.sortie).exists():
                print("%s existe déjà : choisir un autre nom." % a.sortie)
                return 1
            corrige = lire_corrige(a.corrige)
            resultat = consolider(lire_campagne(a.campagne), corrige)
            ecrire_synthese(resultat, corrige, a.sortie)
            _afficher(resultat)
            print("\nSynthèse écrite : %s" % a.sortie)
    except ErreurEvaluation as e:
        print("À corriger avant de continuer (%d) :" % len(e.erreurs))
        for err in e.erreurs:
            print("  - " + err)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(principal())
