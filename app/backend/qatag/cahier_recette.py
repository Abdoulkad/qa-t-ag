"""Cahier de recette Word (RA-32, RA-33), au gabarit DSI du rapport de test.

Le gabarit (`webapp/modele-rapport-test.docx`) est rempli depuis la base, sans
reformulation : page de garde, contexte de la mission, objectifs (critères
d'acceptation validés à PV-1), un bloc par cas avec son dernier passage
(résultat observé, statut coché, captures, testeur et date, commentaires et
anomalies), tableau récapitulatif, conclusion (compteurs, anomalies, décision
de clôture et levée des réserves). Le fichier produit est ensuite relu : chaque
cas doit y porter le statut enregistré en base, sinon rien n'est remis.
"""
import copy
import hashlib
import io
import json
import os
import tempfile

import docx
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from sqlalchemy import select
from sqlalchemy.orm import Session, undefer

from . import journal
from . import rapport as rapport_svc
from .config import parametres
from .missions import lignes_cahier, version_courante
from .models import (Anomalie, ClotureRecette, CritereAcceptation, DecisionPorte, Execution,
                     Mission, Organisation, Preuve, Utilisateur)
from .recette import dernieres_executions
from .regles import LIBELLES_TEST, RegleViolee

STATUTS_EXPORT = ("EN_RECETTE", "CLOTUREE", "RESERVES_LEVEES")
ORGANISATION_PAR_DEFAUT = "Direction du Système d'Information"
# Statut en base -> libellé de la case à cocher du gabarit
CASES = {"PASS": "PASS", "FAIL": "FAIL", "BLOCKED": "BLOQUÉ", "NOT_RUN": "NON EXÉCUTÉ"}
TYPES = {"Fonctionnel": "Fonctionnel", "Sécurité/Habilitation": "Sécurité / Habilitation",
         "Interface utilisateur": "Interface utilisateur"}
LIBELLES_ANOMALIE = {"OUVERTE": "ouverte", "CORRIGEE": "corrigée, à retester",
                     "FERMEE": "fermée", "REJETEE": "rejetée"}


class ControleEchoue(Exception):
    """RA-33 : le fichier produit ne correspond pas à la base."""


def _nom(db: Session, uid: int) -> str:
    p = db.get(Utilisateur, uid)
    return (p.nom or p.email) if p else "?"


# --- Outils Word ------------------------------------------------------------------------

def _ligne(table, debut: str):
    """Ligne du bloc de cas dont la première cellule commence par `debut`."""
    for row in table.rows:
        if row.cells[0].text.strip().startswith(debut):
            return row
    raise ControleEchoue("gabarit inattendu : ligne « %s » absente" % debut)


def _texte_normal(run, texte: str) -> None:
    """Remplace le texte d'un run d'aide (italique gris) par du contenu réel."""
    run.text = texte
    run.italic = False
    run.font.color.rgb = None


def _ecrire(cellule, texte: str) -> None:
    """Contenu réel dans une cellule : premier paragraphe, paragraphes vides ensuite
    supprimés au-delà du nécessaire, mise en forme d'aide retirée."""
    lignes = (texte or "—").split("\n")
    paragraphes = cellule.paragraphs
    for i, p in enumerate(paragraphes):
        valeur = lignes[i] if i < len(lignes) else ""
        if p.runs:
            _texte_normal(p.runs[0], valeur)
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
        elif valeur:
            p.add_run(valeur)
    for reste in lignes[len(paragraphes):]:
        cellule.add_paragraph(reste)
    # Les lignes vides du gabarit, inutiles une fois rempli
    for p in cellule.paragraphs[len(lignes):]:
        if not p.text.strip() and len(cellule.paragraphs) > 1:
            p._element.getparent().remove(p._element)


def _cocher(cellule, libelle: str) -> None:
    """☒ devant le libellé choisi, en gardant la mise en forme du gabarit."""
    for p in cellule.paragraphs:
        for r in p.runs:
            if r.text.strip() == "☐  %s" % libelle or r.text.strip() == "☐ %s" % libelle:
                r.text = r.text.replace("☐", "☒", 1)
                return


def _apres(ancre, texte: str, modele, gras: bool = False):
    """Nouveau paragraphe inséré après `ancre` (élément), au style de `modele`."""
    p = OxmlElement("w:p")
    ancre.addnext(p)
    paragraphe = Paragraph(p, modele._parent)
    if modele.style is not None:
        paragraphe.style = modele.style
    if texte:
        paragraphe.add_run(texte).bold = gras
    return paragraphe


def _puce(ancre, texte: str, modele):
    """Puce en retrait, au style du paragraphe modèle (le gabarit n'a pas de liste à puces)."""
    p = _apres(ancre, "", modele)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Cm(0.1)
    p.add_run("•  ")
    p.add_run(texte)
    return p


def _bloc_texte(ancre, texte: str, modele):
    """Paragraphes d'un texte rédigé ; les lignes « - … » deviennent des puces."""
    for ligne in (x.strip() for x in texte.split("\n")):
        if not ligne:
            continue
        if ligne[:2] in ("- ", "* ", "• "):
            ancre = _puce(ancre, ligne[2:].strip(), modele)._element
        else:
            ancre = _apres(ancre, ligne, modele)._element
    return ancre


def _note(ancre, texte: str, modele):
    """Mention discrète (traçabilité, validation) : petite, en italique, grise."""
    p = _apres(ancre, "", modele)
    run = p.add_run(texte)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return p


def _paragraphe(doc, debut: str, style: str | None = None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(debut) and (style is None or p.style.name == style):
            return p
    raise ControleEchoue("gabarit inattendu : paragraphe « %s » absent" % debut)


def _blocs(doc) -> list[tuple[str, object]]:
    """(code du cas, table) pour chaque bloc de cas, dans l'ordre du document."""
    blocs = []
    for t in doc.tables:
        if not any(r.cells[0].text.strip().startswith("Statut d'exécution") for r in t.rows):
            continue
        precedent = t._element.getprevious()
        texte = "".join(x.text or "" for x in precedent.iter() if x.tag.endswith("}t")) \
            if precedent is not None else ""
        blocs.append((texte.split(" : ")[0].strip(), t))
    return blocs


def _statut_coche(table) -> str | None:
    texte = _ligne(table, "Statut d'exécution").cells[-1].text
    coches = [code for code, libelle in CASES.items() if "☒  %s" % libelle in texte]
    return coches[0] if len(coches) == 1 else None


NOMBRES = ("zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf",
           "dix", "onze", "douze", "treize", "quatorze", "quinze", "seize", "dix-sept",
           "dix-huit", "dix-neuf", "vingt")


def _en_lettres(n: int) -> str:
    return NOMBRES[n] if 0 <= n < len(NOMBRES) else str(n)


def _part(n: int, total: int) -> str:
    """« 03 (75 %) », comme dans les rapports de test de la DSI."""
    return "%02d (%d %%)" % (n, round(100 * n / total)) if total else "00"


def _recapitulatif(doc):
    for t in doc.tables:
        if len(t.columns) >= 5 and t.rows and "RESULTAT OBTENU" in t.rows[0].cells[4].text:
            return t
    raise ControleEchoue("gabarit inattendu : tableau récapitulatif absent")


# --- Export ------------------------------------------------------------------------------

def exporter(db: Session, u: Utilisateur, m: Mission) -> tuple[bytes, str]:
    if m.statut not in STATUTS_EXPORT:
        raise RegleViolee("RA-32", "Le cahier de recette s'exporte une fois la recette "
                                   "commencée (statut actuel : %s)" % m.statut)
    v = version_courante(m)
    cas, lignes = lignes_cahier(db, v, "RA-32")
    gabarit = {ligne[0]: (ligne[5], ligne[6]) for ligne in lignes}      # priorité, type
    ids = [c.id for c in cas]
    dernieres = dernieres_executions(db, ids)
    passages = {c.id: 0 for c in cas}
    dates = []
    testeurs = set()
    for e in db.scalars(select(Execution).where(Execution.cas_id.in_(ids))):
        passages[e.cas_id] += 1
        dates.append(e.execute_le)
        testeurs.add(e.testeur_id)
    preuves: dict[int, list[Preuve]] = {}
    for p in db.scalars(select(Preuve).options(undefer(Preuve.contenu)).where(
            Preuve.execution_id.in_([e.id for e in dernieres.values()])).order_by(Preuve.id)):
        preuves.setdefault(p.execution_id, []).append(p)
    anomalies = db.scalars(select(Anomalie).where(Anomalie.mission_id == m.id)
                           .order_by(Anomalie.code)).all()
    criteres = db.scalars(select(CritereAcceptation).where(CritereAcceptation.version_id == v.id)
                          .order_by(CritereAcceptation.code)).all()
    cloture = db.scalar(select(ClotureRecette).where(ClotureRecette.mission_id == m.id))
    pv2 = db.scalar(select(DecisionPorte).where(DecisionPorte.version_id == v.id,
                                                DecisionPorte.porte == "PV-2",
                                                DecisionPorte.decision == "APPROUVE"))
    org = db.get(Organisation, m.organisation_id)
    nom_org = org.nom if org and org.nom and org.nom != "Mon organisation" else ORGANISATION_PAR_DEFAUT

    compte = {"PASS": 0, "FAIL": 0, "BLOCKED": 0, "NOT_RUN": 0, "aucun": 0}
    for c in cas:
        e = dernieres.get(c.id)
        compte[e.statut if e else "aucun"] += 1

    # RA-68 : parties rédigées validées, recontrôlées contre les faits du moment
    redaction, rapport_obj = rapport_svc.pour_export(db, m)

    doc = docx.Document(str(parametres().gabarit_rapport_test))
    normal = _paragraphe(doc, "Les objectifs de cette séance")
    # Titres et puces vides du gabarit : ils apparaîtraient comme des numéros sans texte.
    # Un paragraphe qui porte un saut de section est gardé (mise en page), en style normal.
    for p in list(doc.paragraphs):
        ppr = p._p.pPr
        numerote = ppr is not None and ppr.numPr is not None
        if not p.text.strip() and (p.style.name.startswith("Heading") or numerote):
            if ppr is not None and ppr.sectPr is not None:
                p.style = normal.style
                if ppr.numPr is not None:
                    ppr.remove(ppr.numPr)
            else:
                p._element.getparent().remove(p._element)

    # Page de garde
    for p in doc.paragraphs:
        if "[NOM DE L'ORGANISATION]" in p.text:
            for r in p.runs:
                r.text = r.text.replace("[NOM DE L'ORGANISATION]", nom_org)

    # Contexte : le récit (rédigé et validé, sinon une présentation sobre), puis une
    # note de traçabilité
    rejoues = sum(1 for c in cas if passages[c.id] > 1)
    periode = ("le %s" % min(dates).strftime("%d/%m/%Y") if dates and
               min(dates).date() == max(dates).date() else
               "du %s au %s" % (min(dates).strftime("%d/%m/%Y"), max(dates).strftime("%d/%m/%Y"))
               if dates else "")
    ancre = _paragraphe(doc, "Contexte", "Heading 1")._element
    if not cloture:
        ancre = _apres(ancre, "DOCUMENT DE TRAVAIL — recette en cours, non clôturée.",
                       normal, gras=True)._element
    if redaction:
        ancre = _bloc_texte(ancre, redaction["contexte"], normal)
    else:
        recit = ("Le présent document constitue le rapport de test de la fonctionnalité "
                 "« %s ». Il rend compte de l'exécution des %s (%02d) cas de test validés "
                 "à la Human Gate PV-2%s." % (m.titre, _en_lettres(len(cas)), len(cas),
                                             ", menée %s" % periode if periode else ""))
        if rejoues:
            recit += (" %s (%02d) cas %s été rejoué%s après la correction d'anomalies : le "
                      "rapport vaut aussi rapport de tests de confirmation." % (
                          _en_lettres(rejoues).capitalize(), rejoues,
                          "ont" if rejoues > 1 else "a", "s" if rejoues > 1 else ""))
        ancre = _apres(ancre, recit, normal)._element
    reference = "Références : mission %s, version %d de l'exigence (empreinte SHA-256 %s)" % (
        m.reference, v.numero, v.empreinte[:16])
    if pv2:
        reference += ", cas validés à PV-2 le %s par %s%s" % (
            pv2.decide_le.strftime("%d/%m/%Y"), _nom(db, pv2.decideur_id),
            " (auto-validation)" if pv2.auto_validation else "")
    if periode:
        reference += ", exécution %s" % periode
    reference += " ; testeur%s : %s." % ("s" if len(testeurs) > 1 else "",
                                        ", ".join(sorted(_nom(db, t) for t in testeurs)) or "—")
    _note(ancre, reference, normal)

    # Objectifs : ceux du rapport validé, sinon un objectif par critère validé à PV-1
    ancre = normal._element
    objectifs = (redaction["objectifs"] if redaction else
                 ["Vérifier le critère %s : %s" % (c.code, c.texte.rstrip(".")) for c in criteres]) \
        or ["Vérifier les cas de test validés à PV-2."]
    for objectif in objectifs:
        ancre = _puce(ancre, objectif.rstrip(".") + ".", normal)._element

    # Exécution des tests : un bloc par cas, cloné depuis le modèle du gabarit
    titre_modele = _paragraphe(doc, "[CT-XXX]")
    blocs = _blocs(doc)
    if len(blocs) != 1:
        raise ControleEchoue("gabarit inattendu : un seul bloc de cas attendu")
    table_modele = blocs[0][1]
    for consigne in ("Modèle de cas de test", "Dupliquer ce bloc"):
        p = _paragraphe(doc, consigne)
        p._element.getparent().remove(p._element)
    p_original = copy.deepcopy(titre_modele._element)
    t_original = copy.deepcopy(table_modele._element)
    ancre = table_modele._element
    elements = [(titre_modele._element, table_modele._element)]
    for _ in cas[1:]:
        nouveau_p, nouvelle_t = copy.deepcopy(p_original), copy.deepcopy(t_original)
        ancre.addnext(nouveau_p)
        nouveau_p.addnext(nouvelle_t)
        elements.append((nouveau_p, nouvelle_t))
        ancre = nouvelle_t
    for c, (p_el, t_el) in zip(cas, elements):
        titre = Paragraph(p_el, titre_modele._parent)
        titre.runs[0].text = "%s : " % c.code
        titre.runs[1].text = c.titre
        for r in titre.runs[2:]:
            r._element.getparent().remove(r._element)
        table = Table(t_el, table_modele._parent)
        e = dernieres.get(c.id)

        description = table.rows[1].cells[-1]
        pars = description.paragraphs
        _texte_normal(pars[0].runs[1], c.objectif or "—")
        _texte_normal(pars[1].runs[1], c.donnees_test or "—")
        if c.preconditions:
            prec = _apres(pars[0]._element, "", pars[0])
            prec.add_run("Préconditions : ").bold = True
            prec.add_run(c.preconditions)
        etapes = [x.strip() for x in (c.etapes or "").split("\n") if x.strip()] or ["—"]
        lignes_etapes = description.paragraphs[4 if c.preconditions else 3:]
        modele_etape = lignes_etapes[0]
        for p in lignes_etapes[1:]:
            p._element.getparent().remove(p._element)
        _texte_normal(modele_etape.runs[0], etapes[0])
        ancre_etape = modele_etape._element
        for etape in etapes[1:]:
            ancre_etape = _apres(ancre_etape, etape, modele_etape)._element

        _ecrire(_ligne(table, "Règle / Exigence").cells[-1], c.exigence)
        priorite, type_ = gabarit[c.code]
        _cocher(_ligne(table, "Priorité").cells[-1], priorite)
        _cocher(_ligne(table, "Type de test").cells[-1], TYPES.get(type_, type_))
        _ecrire(_ligne(table, "Résultat attendu").cells[-1], c.resultat_attendu)
        _cocher(_ligne(table, "Statut d'exécution").cells[-1], CASES[e.statut if e else "NOT_RUN"])
        observe = _ligne(table, "Résultat observé").cells[-1]
        capture = _ligne(table, "Capture d'écran").cells[-1]
        testeur = _ligne(table, "Testeur & Date").cells[-1]
        commentaire = []
        if e is None:
            _ecrire(observe, "Cas non exécuté.")
            _ecrire(capture, "—")
            _ecrire(testeur, "Nom du testeur : —          Date : —")
        else:
            _ecrire(observe, "Passage %d du %s\n%s" % (
                passages[c.id], e.execute_le.strftime("%d/%m/%Y %H:%M"),
                e.resultat_obtenu or "—"))
            images = preuves.get(e.id, [])
            _ecrire(capture, "" if images else "Aucune capture.")
            for p in images:
                capture.paragraphs[-1].add_run().add_picture(io.BytesIO(p.contenu), width=Cm(12))
                capture.add_paragraph("%s · SHA-256 %s" % (p.nom, p.empreinte[:16]))
            _ecrire(testeur, "Nom du testeur : %s          Date : %s" % (
                _nom(db, e.testeur_id), e.execute_le.strftime("%d / %m / %Y")))
            if e.commentaire:
                commentaire.append(e.commentaire)
        commentaire += ["%s (%s, %s) : %s" % (a.code, a.gravite or "à qualifier",
                                              LIBELLES_ANOMALIE[a.statut], a.titre)
                        for a in anomalies if a.cas_id == c.id]
        _ecrire(_ligne(table, "Commentaires").cells[-1], "\n".join(commentaire) or "—")

    # Récapitulatif : comptage, puis une ligne par cas (porteur de l'action à renseigner)
    recap = _recapitulatif(doc)
    executes = compte["PASS"] + compte["FAIL"] + compte["BLOCKED"]
    non_executes = compte["NOT_RUN"] + compte["aucun"]
    phrase = recap._element.getprevious()
    phrase_p = _apres(phrase, "Sur les %s (%02d) cas de test, %s (%02d) %s été exécuté%s." % (
        _en_lettres(len(cas)), len(cas), _en_lettres(executes), executes,
        "ont" if executes > 1 else "a", "s" if executes > 1 else ""), normal)
    comptage = doc.add_table(rows=2, cols=6)
    if "Table Grid" in [st.name for st in doc.styles]:
        comptage.style = doc.styles["Table Grid"]
    entetes = ("FONCTIONNALITÉ", "SUCCÈS", "ÉCHEC", "NON TESTABLE", "NON EXÉCUTÉ", "TOTAL")
    valeurs = (m.titre, _part(compte["PASS"], len(cas)), _part(compte["FAIL"], len(cas)),
               _part(compte["BLOCKED"], len(cas)), _part(non_executes, len(cas)),
               "%02d" % len(cas))
    for cellule, texte in zip(comptage.rows[0].cells, entetes):
        cellule.text = ""
        cellule.paragraphs[0].add_run(texte).bold = True
    for cellule, texte in zip(comptage.rows[1].cells, valeurs):
        cellule.text = texte
    comptage.autofit = False
    for ligne in comptage.rows:
        for i, cellule in enumerate(ligne.cells):
            cellule.width = Cm(7) if i == 0 else Cm(3.4)
    phrase_p._element.addnext(comptage._tbl)
    _apres(comptage._tbl, "", normal)
    ligne_modele = copy.deepcopy(recap.rows[1]._tr)
    for tr in [r._tr for r in recap.rows[1:]]:
        tr.getparent().remove(tr)
    for i, c in enumerate(cas, start=1):
        tr = copy.deepcopy(ligne_modele)
        recap._tbl.append(tr)
        cellules = recap.rows[-1].cells
        e = dernieres.get(c.id)
        du_cas = [a for a in anomalies if a.cas_id == c.id]
        obtenu = ("%s — %s" % (LIBELLES_TEST[e.statut], e.resultat_obtenu) if e and e.resultat_obtenu
                  else LIBELLES_TEST[e.statut] if e else "Non exécuté")
        if du_cas:
            impact = "\n".join("%s : %s" % (a.code, a.gravite or "à qualifier") for a in du_cas)
        elif e and e.statut == "PASS":
            impact = "Aucun"
        else:
            impact = "Cas non vérifié"
        observations = []
        if redaction and redaction["recommandations"].get(c.code):
            observations.append(redaction["recommandations"][c.code])
        observations += [e.commentaire] if e and e.commentaire else []
        observations += ["%s %s : %s" % (a.code, LIBELLES_ANOMALIE[a.statut], a.titre)
                         for a in du_cas]
        for cellule, valeur in zip(cellules, (
                str(i), "%s : %s" % (c.code, c.titre), c.objectif, c.resultat_attendu, obtenu,
                impact, "\n".join(observations), "")):
            _ecrire(cellule, valeur) if valeur else _ecrire(cellule, "")

    # Synthèse (section ajoutée avant la conclusion, comme dans les rapports de la DSI)
    titre_conclusion = _paragraphe(doc, "Conclusion", "Heading 1")
    titre_synthese = copy.deepcopy(titre_conclusion._element)
    titre_conclusion._element.addprevious(titre_synthese)
    p_synthese = Paragraph(titre_synthese, titre_conclusion._parent)
    p_synthese.runs[0].text = "Synthèse"
    for run in p_synthese.runs[1:]:
        run._element.getparent().remove(run._element)
    ancre = titre_synthese
    if redaction:
        _bloc_texte(ancre, redaction["synthese"], normal)
    else:
        resultat = ("Les %s (%02d) cas de test exécutés sont conformes aux résultats attendus."
                    % (_en_lettres(len(cas)), len(cas))
                    if compte["PASS"] == len(cas) else
                    "Sur les %s (%02d) cas de test, %s (%02d) %s conforme%s aux résultats attendus"
                    % (_en_lettres(len(cas)), len(cas), _en_lettres(compte["PASS"]),
                       compte["PASS"], "sont" if compte["PASS"] > 1 else "est",
                       "s" if compte["PASS"] > 1 else "") +
                    "".join(", %s (%02d) %s" % (_en_lettres(n), n, libelle) for n, libelle in (
                        (compte["FAIL"], "en échec"), (compte["BLOCKED"], "non testable%s"
                                                       % ("s" if compte["BLOCKED"] > 1 else "")),
                        (non_executes, "non exécuté%s" % ("s" if non_executes > 1 else "")))
                        if n) + ".")
        ancre = _apres(ancre, "La fonctionnalité « %s » a fait l'objet de tests fonctionnels "
                              "visant à vérifier les points suivants :" % m.titre,
                       normal)._element
        points = (["le critère %s : %s" % (c.code, c.texte.rstrip(".")) for c in criteres]
                  or ["les cas de test validés à PV-2"])
        for i, point in enumerate(points):
            ancre = _puce(ancre, point + (" ;" if i < len(points) - 1 else "."), normal)._element
        ancre = _apres(ancre, resultat, normal)._element
        ouvertes = [a for a in anomalies if a.statut in ("OUVERTE", "CORRIGEE")]
        if ouvertes:
            _apres(ancre, "Anomalie%s restant à traiter : %s." % (
                "s" if len(ouvertes) > 1 else "",
                " ; ".join("%s (%s, %s) — %s" % (a.code, (a.gravite or "à qualifier").lower(),
                                                  LIBELLES_ANOMALIE[a.statut], a.titre)
                           for a in ouvertes)), normal)
        elif anomalies:
            _apres(ancre, "Les anomalies relevées (%s) sont toutes résolues." %
                   ", ".join(a.code for a in anomalies), normal)
        else:
            _apres(ancre, "Aucune anomalie n'a été constatée au cours de cette phase de tests.",
                   normal)

    # Conclusion : texte validé et avis, puis la décision de clôture et la validation
    ancre = titre_conclusion._element
    if redaction:
        ancre = _bloc_texte(ancre, redaction["conclusion"], normal)
        ancre = _apres(ancre, "Avis sur la mise en production : %s." % redaction["avis"],
                       normal, gras=True)._element
    if cloture:
        decision = "La recette a été clôturée %s le %s par %s%s" % (
            "avec réserves" if cloture.decision == "CLOTURE_AVEC_RESERVES" else "sans réserve",
            cloture.decide_le.strftime("%d/%m/%Y"), _nom(db, cloture.decideur_id),
            " (auto-validation)" if cloture.auto_validation else "")
        decision += (", au motif suivant : « %s »." % cloture.commentaire.rstrip(".")
                     if cloture.commentaire else ".")
        ancre = _apres(ancre, decision, normal)._element
        reserves = json.loads(cloture.reserves)
        if reserves:
            ancre = _apres(ancre, "Les réserves portent sur les points suivants :", normal)._element
            for reserve in reserves:
                ancre = _puce(ancre, reserve, normal)._element
        if cloture.levee_le:
            levee = "Ces réserves ont été levées le %s par %s%s" % (
                cloture.levee_le.strftime("%d/%m/%Y"), _nom(db, cloture.levee_par_id),
                " (auto-validation)" if cloture.levee_auto_validation else "")
            levee += (" : %s" % cloture.levee_commentaire.rstrip(".") + "."
                      if cloture.levee_commentaire else ".")
            ancre = _apres(ancre, levee, normal)._element
        elif reserves:
            ancre = _apres(ancre, "Ces réserves sont en cours de levée.", normal)._element
    elif not redaction:
        ancre = _apres(ancre, "La recette est en cours : les conclusions seront arrêtées à sa "
                              "clôture.", normal)._element
    if redaction:
        auteur = ("Parties rédigées proposées par l'agent AG-11, relues et validées par %s le %s."
                  if rapport_obj.proposition else "Parties rédigées validées par %s le %s.")
        _note(ancre, auteur % (_nom(db, rapport_obj.valide_par_id),
                               rapport_obj.valide_le.strftime("%d/%m/%Y")), normal)

    fd, chemin = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        doc.save(chemin)
        # RA-33 : relecture du fichier produit, comparaison à la base
        relu = docx.Document(chemin)
        blocs_relus = _blocs(relu)
        if [code for code, _ in blocs_relus] != [c.code for c in cas]:
            raise ControleEchoue("les cas du document ne correspondent pas à la base")
        for c, (_, table) in zip(cas, blocs_relus):
            attendu = dernieres[c.id].statut if c.id in dernieres else "NOT_RUN"
            if _statut_coche(table) != attendu:
                raise ControleEchoue("statut de %s différent de la base" % c.code)
        lignes_recap = _recapitulatif(relu).rows[1:]
        if [r.cells[1].text.split(" : ")[0] for r in lignes_recap] != [c.code for c in cas]:
            raise ControleEchoue("le tableau récapitulatif ne correspond pas aux cas")
        texte = "\n".join(p.text for p in relu.paragraphs) + "\n".join(
            cell.text for t in relu.tables for row in t.rows for cell in row.cells)
        absentes = [a.code for a in anomalies if a.code not in texte]
        if absentes:
            raise ControleEchoue("anomalies absentes du document : %s" % ", ".join(absentes))
        if "[CT-XXX]" in texte or "[Décrire" in texte:
            raise ControleEchoue("zones du gabarit restées vides")
        with open(chemin, "rb") as f:
            contenu = f.read()
    finally:
        os.remove(chemin)
    nom = "Cahier_recette_%s_v%d%s.docx" % (m.reference, v.numero, "" if cloture else "_brouillon")
    journal.ajouter(db, m.organisation_id, "cahier_recette_exporte", mission_id=m.id,
                    acteur_id=u.id, details={
                        "version": v.numero, "cas": len(cas), "anomalies": len(anomalies),
                        "cloturee": bool(cloture), "fichier": nom, "controle": "conforme",
                        "gabarit": "rapport de test",
                        "rapport_redige": bool(redaction),
                        "empreinte": hashlib.sha256(contenu).hexdigest()})
    db.commit()
    return contenu, nom
