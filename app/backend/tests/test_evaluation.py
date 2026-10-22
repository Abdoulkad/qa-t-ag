"""Outillage de la campagne d'évaluation des agents : export, appariement, indicateurs."""
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from qatag import evaluation as ev
from qatag.db import SessionLocale

from .conftest import US_TEST


# --- Corrigé et campagne simulés ----------------------------------------------------------

def _remplir(chemin, feuille, lignes):
    wb = load_workbook(chemin)
    ws = wb[feuille]
    entete = [c.value for c in ws[1]]
    for ligne in lignes:
        ws.append([ligne.get(c) for c in entete])
    wb.save(chemin)


def _modifier(chemin, feuille, cle, valeurs):
    """Saisie du relecteur : valeurs posées sur la ligne dont les colonnes valent `cle`."""
    wb = load_workbook(chemin)
    ws = wb[feuille]
    entete = [c.value for c in ws[1]]
    trouve = False
    for row in ws.iter_rows(min_row=2):
        ligne = dict(zip(entete, (c.value for c in row)))
        if all(str(ligne.get(k)) == str(v) for k, v in cle.items()):
            for col, val in valeurs.items():
                row[entete.index(col)].value = val
            trouve = True
    wb.save(chemin)
    assert trouve, cle


def _corrige(tmp_path, signe=True):
    chemin = tmp_path / "corrige.xlsx"
    ev.modele_corrige(chemin)
    _remplir(chemin, "Corpus", [{"Exigence": "EX-01", "Type": "User Story courte",
                                 "Titre": "Réinitialisation du mot de passe"}])
    _remplir(chemin, "Défauts", [
        {"Exigence": "EX-01", "VT": "VT-01", "Citation": "durée limitée",
         "Catégorie": "Ambiguïté", "Criticité": "Bloquante"},
        {"Exigence": "EX-01", "VT": "VT-02", "Citation": "sécurisé",
         "Catégorie": "Ambiguïté", "Criticité": "Majeure"},
        {"Exigence": "EX-01", "VT": "VT-03", "Citation": "REQ-001",
         "Catégorie": "Règle manquante", "Criticité": "Bloquante"},
    ])
    _remplir(chemin, "Règles", [
        {"Exigence": "EX-01", "Règle": "RC-01", "Énoncé": "Le lien est valide une durée limitée"},
        {"Exigence": "EX-01", "Règle": "RC-02", "Énoncé": "Le nouveau mot de passe est sécurisé"},
        {"Exigence": "EX-01", "Règle": "RC-03", "Énoncé": "Un e-mail confirme le changement"},
    ])
    if signe:
        wb = load_workbook(chemin)
        for row in wb["Signature"].iter_rows(min_row=2):
            row[1].value = "2026-10-22" if "Date" in row[0].value else "Relecteur"
        wb.save(chemin)
    return chemin


def _mission_eval(client, h, ia, titre):
    """Mission menée jusqu'à la conception (PV-2 en attente), titre de campagne."""
    r = client.post("/api/missions", json={"titre": titre, "texte": US_TEST}, headers=h)
    assert r.status_code == 201, r.text
    mid = r.json()["mission"]["id"]
    v = client.get("/api/missions/%d" % mid, headers=h).json()["versions"][0]
    for q, rep in zip(v["questions"], ("30 minutes", "8 caractères minimum")):
        client.put("/api/missions/%d/questions/%d" % (mid, q["id"]), json={"reponse": rep},
                   headers=h)
    ia.reponses["reecriture"] = ia.reponses["reecriture"].split("### CA-03")[0]
    client.post("/api/missions/%d/reecriture" % mid, headers=h)
    r = client.post("/api/missions/%d/portes/PV-1" % mid, json={"decision": "APPROUVE"}, headers=h)
    assert r.status_code == 200, r.text
    return mid


def _exporter(tmp_path, corrige, nom="campagne.xlsx", reprendre=None):
    with SessionLocale() as db:
        missions = ev.extraire(db)
    chemin = tmp_path / nom
    ev.ecrire_campagne(missions, ev.lire_corrige(corrige), chemin, reprendre)
    return chemin, missions


def _saisir_campagne(chemin):
    """P1 : DEF-003 jugé trouvaille ; P2 : jugé FP et RG-02 inventée."""
    p1, p2 = "EVAL-EX-01-DeepSeek-P1", "EVAL-EX-01-DeepSeek-P2"
    for m in (p1, p2):
        _modifier(chemin, "Défauts", {"Mission": m, "Code": "DEF-001"}, {"Décision": "VP", "VT": "VT-01"})
        _modifier(chemin, "Défauts", {"Mission": m, "Code": "DEF-002"}, {"Décision": "VP", "VT": "VT-02"})
        _modifier(chemin, "Règles", {"Mission": m, "Code": "RG-01"}, {"Règle du corrigé": "RC-01"})
        for cas in ("CT-001", "CT-002", "CT-003", "CT-004"):
            _modifier(chemin, "Cas", {"Mission": m, "Cas": cas},
                      {c: "oui" for c in ev.CRITERES_CAS})
    _modifier(chemin, "Défauts", {"Mission": p1, "Code": "DEF-003"}, {"Décision": "Trouvaille"})
    _modifier(chemin, "Défauts", {"Mission": p2, "Code": "DEF-003"}, {"Décision": "FP"})
    _modifier(chemin, "Règles", {"Mission": p1, "Code": "RG-02"}, {"Règle du corrigé": "RC-02"})
    _modifier(chemin, "Règles", {"Mission": p2, "Code": "RG-02"}, {"Inventée": "oui"})
    _modifier(chemin, "Cas", {"Mission": p1, "Cas": "CT-004"}, {"Données exécutables": "non"})


# --- Noms, similarité, appariement ----------------------------------------------------

def test_nom_de_mission_de_campagne():
    assert ev.lire_nom("EVAL-EX-01-DeepSeek-P2") == ("EX-01", "DeepSeek", 2)
    assert ev.lire_nom("EVAL-EX-12-Ollama-P3 relance") == ("EX-12", "Ollama", 3)
    assert ev.lire_nom("Réinitialisation du mot de passe") is None


def test_similarite_des_citations():
    assert ev.similarite("« durée limitée »", "Durée limitée") == 1.0
    assert ev.similarite("valable pendant une durée limitée", "durée limitée") == 1.0
    assert ev.similarite("mot de passe sécurisé", "sans effort") == 0.0


def test_appariement_un_vp_par_defaut_du_corrige_puis_doublon_puis_fp():
    produits = [{"citation": "durée limitée"}, {"citation": "une durée limitée"},
                {"citation": "facilement"}]
    corrige = [{"vt": "VT-01", "citation": "durée limitée"}]
    assert ev.apparier(produits, corrige, "citation", "citation", "vt") == [
        ("VP", "VT-01", 1.0), ("Doublon", "VT-01", 1.0), ("FP", "", 0.0)]


def test_passage_echoue_seulement_apres_deux_echecs_ou_echec_final():
    t = lambda i, type_, etat: SimpleNamespace(id=i, type=type_, etat=etat)  # noqa: E731
    assert not ev._passage_echoue([t(1, "ANALYSE", "ECHEC"), t(2, "ANALYSE", "TERMINEE")])
    assert ev._passage_echoue([t(1, "ANALYSE", "ECHEC"), t(2, "ANALYSE", "ECHEC"),
                               t(3, "ANALYSE", "TERMINEE")])
    assert ev._passage_echoue([t(1, "ANALYSE", "TERMINEE"), t(2, "CONCEPTION", "ECHEC")])


# --- Corrigé -----------------------------------------------------------------------------

def test_corrige_vierge_refuse_puis_avertit_s_il_n_est_pas_signe(tmp_path):
    vierge = tmp_path / "vierge.xlsx"
    ev.modele_corrige(vierge)
    with pytest.raises(ev.ErreurEvaluation, match="aucune exigence"):
        ev.lire_corrige(vierge)
    corrige = ev.lire_corrige(_corrige(tmp_path, signe=False))
    assert corrige["seuils"]["I1"] == ("Rappel (%)", ">=", 80.0, "Éliminatoire")
    assert any("non signé" in a for a in corrige["avertissements"])


def test_corrige_refuse_criticite_hors_liste_et_vt_en_double(tmp_path):
    chemin = _corrige(tmp_path)
    _remplir(chemin, "Défauts", [{"Exigence": "EX-01", "VT": "VT-01", "Citation": "x",
                                  "Catégorie": "Ambiguïté", "Criticité": "Grave"}])
    with pytest.raises(ev.ErreurEvaluation) as e:
        ev.lire_corrige(chemin)
    assert any("VT-01 en double" in m for m in e.value.erreurs)
    assert any("Criticité « Grave » hors liste" in m for m in e.value.erreurs)


# --- Export depuis la base -----------------------------------------------------------------

def test_export_des_sorties_brutes_avec_propositions(client, admin, ia, tmp_path):
    _mission_eval(client, admin, ia, "EVAL-EX-01-DeepSeek-P1")
    client.post("/api/missions", json={"titre": "Mission de travail", "texte": US_TEST},
                headers=admin)
    chemin, missions = _exporter(tmp_path, _corrige(tmp_path))
    assert [m["mission"] for m in missions] == ["EVAL-EX-01-DeepSeek-P1"]
    m = missions[0]
    assert (m["exigence"], m["config"], m["passage"]) == ("EX-01", "DeepSeek", 1)
    assert [d["code"] for d in m["defauts"]] == ["DEF-001", "DEF-002", "DEF-003"]
    assert len(m["cas"]) == 4 and m["couverture"] is not None and not m["echoue"]

    camp = ev.lire_campagne(chemin)
    props = {l["Code"]: (l["Proposition"], l["VT proposé"]) for l in camp["Défauts"]}
    assert props == {"DEF-001": ("VP", "VT-01"), "DEF-002": ("VP", "VT-02"), "DEF-003": ("FP", "")}
    assert all(l.get("Décision", "") == "" for l in camp["Défauts"])   # l'humain décide
    assert {l["Cas"] for l in camp["Cas"]} == {"CT-001", "CT-002", "CT-003", "CT-004"}
    assert camp["Missions"][0]["Passage échoué"] == "non"


def test_export_refuse_une_exigence_absente_du_corpus(client, admin, ia, tmp_path):
    _mission_eval(client, admin, ia, "EVAL-EX-09-DeepSeek-P1")
    with pytest.raises(ev.ErreurEvaluation, match="EX-09"):
        _exporter(tmp_path, _corrige(tmp_path))


def test_nouvel_export_reprend_les_saisies(client, admin, ia, tmp_path):
    _mission_eval(client, admin, ia, "EVAL-EX-01-DeepSeek-P1")
    corrige = _corrige(tmp_path)
    premier, _ = _exporter(tmp_path, corrige)
    _modifier(premier, "Défauts", {"Code": "DEF-003"}, {"Décision": "Trouvaille",
                                                       "Commentaire": "Vu par les deux"})
    _mission_eval(client, admin, ia, "EVAL-EX-01-DeepSeek-P2")
    second, missions = _exporter(tmp_path, corrige, "campagne-2.xlsx", reprendre=premier)
    assert len(missions) == 2
    lignes = {(l["Mission"], l["Code"]): l for l in ev.lire_campagne(second)["Défauts"]}
    assert lignes[("EVAL-EX-01-DeepSeek-P1", "DEF-003")]["Décision"] == "Trouvaille"
    assert lignes[("EVAL-EX-01-DeepSeek-P1", "DEF-003")]["Commentaire"] == "Vu par les deux"
    assert lignes[("EVAL-EX-01-DeepSeek-P2", "DEF-003")].get("Décision", "") == ""


# --- Consolidation -----------------------------------------------------------------------

def test_consolidation_refuse_tant_qu_une_decision_manque(client, admin, ia, tmp_path):
    _mission_eval(client, admin, ia, "EVAL-EX-01-DeepSeek-P1")
    corrige = _corrige(tmp_path)
    chemin, _ = _exporter(tmp_path, corrige)
    with pytest.raises(ev.ErreurEvaluation) as e:
        ev.consolider(ev.lire_campagne(chemin), ev.lire_corrige(corrige))
    assert any("DEF-001" in m and "Décision manquante" in m for m in e.value.erreurs)
    assert any("Règle du corrigé ou Inventée" in m for m in e.value.erreurs)
    assert any("quatre conditions" in m for m in e.value.erreurs)


def test_indicateurs_et_decision_sur_une_campagne_complete(client, admin, ia, tmp_path):
    for p in (1, 2):
        _mission_eval(client, admin, ia, "EVAL-EX-01-DeepSeek-P%d" % p)
    corrige = _corrige(tmp_path)
    chemin, _ = _exporter(tmp_path, corrige)
    _saisir_campagne(chemin)
    res = ev.consolider(ev.lire_campagne(chemin), ev.lire_corrige(corrige))
    r = res["configs"]["DeepSeek"]
    c = r["comptes"]
    assert (c["vp"], c["fn"], c["fp"], c["trouvailles"]) == (4, 2, 1, 1)
    i = r["indicateurs"]
    assert i["I1"] == 66.7          # 4 / (4 + 2)
    assert i["I2"] == 2             # VT-03, bloquant, oublié aux deux passages
    assert i["I3"] == 80.0          # 4 / (4 + 1)
    assert i["I4"] == 100.0
    assert i["I5"] == 50.0          # VT-02 : « Terme non défini » au lieu d'« Ambiguïté »
    assert i["I6"] == 66.7          # « facilement et sans effort » absent du texte
    assert i["I7"] == 1
    assert i["I8"] == 1             # RG-02 de P2, jugée inventée et acceptée (SOURCE)
    assert i["I9"] == 50.0          # P1 oublie RC-03, P2 oublie RC-02 et RC-03 : 3 / 6
    assert i["I11"] == 87.5         # 7 cas justes sur 8
    assert i["I12"] == 0.0
    assert i["I13"] == 0.0
    assert not r["verdicts"]["I1"] and not r["verdicts"]["I2"]
    assert r["decision"] == "Non apte"
    assert [b["vt"] for b in r["bloquants"]] == ["VT-03", "VT-03"]
    assert r["confusions"] == [(("Ambiguïté", "Terme non défini"), 2)]
    assert res["avertissements"] == []

    synthese = tmp_path / "synthese.xlsx"
    ev.ecrire_synthese(res, ev.lire_corrige(corrige), synthese)
    wb = load_workbook(synthese)
    lignes = {row[0]: row[1] for row in wb["Synthèse"].iter_rows(values_only=True)}
    assert lignes["Décision"] == "Non apte"
    assert lignes["I1 Rappel (%)"] == "66,7 % ✗"
    assert wb["Bloquants oubliés"].max_row == 3


def _seuils():
    return {k: (v[0], v[1], float(v[2]), v[3]) for k, v in ev.SEUILS.items()}


def _campagne_synthetique(vp_par_mission, fp=0, cas_justes=5, echoue=False):
    """Campagne d'une exigence, trois passages, sans base : pour la règle de décision."""
    corrige = {"exigences": ["EX-01"], "seuils": _seuils(), "avertissements": [],
               "defauts": [{"exigence": "EX-01", "vt": "VT-%02d" % i, "citation": "c%d" % i,
                            "categorie": "Ambiguïté", "criticite": "Majeure"} for i in range(1, 6)],
               "regles": [{"exigence": "EX-01", "regle": "RC-01", "enonce": "r"}]}
    camp = {"Missions": [], "Défauts": [], "Règles": [], "Cas": []}
    n = 2
    for p, vps in enumerate(vp_par_mission, start=1):
        nom = "EVAL-EX-01-X-P%d" % p
        camp["Missions"].append({"Mission": nom, "Exigence": "EX-01", "Config": "X",
                                 "Passage": "P%d" % p, "Passage échoué": "oui" if echoue else "non",
                                 "Couverture %": 100, "_ligne": n})
        for k in range(vps):
            camp["Défauts"].append({"Mission": nom, "Code": "DEF-%d" % k, "Décision": "VP",
                                    "VT": "VT-%02d" % (k + 1), "Catégorie": "Ambiguïté",
                                    "Criticité agent": "Majeure", "Citation": "c",
                                    "Citation vérifiée": "oui", "_ligne": n})
        for k in range(fp):
            camp["Défauts"].append({"Mission": nom, "Code": "FP-%d" % k, "Décision": "FP",
                                    "Citation": "c", "Citation vérifiée": "oui", "_ligne": n})
        camp["Règles"].append({"Mission": nom, "Code": "RG-01", "Règle du corrigé": "RC-01",
                               "_ligne": n})
        for k in range(5):
            juste = "oui" if k < cas_justes else "non"
            camp["Cas"].append({"Mission": nom, "Cas": "CT-%d" % k, "_ligne": n,
                                **{c: juste for c in ev.CRITERES_CAS}})
    return ev.consolider(camp, corrige)["configs"]["X"]


def test_decision_apte_sous_reserve_et_non_apte():
    assert _campagne_synthetique([5, 5, 5])["decision"] == "Apte"
    # Rappel 100 %, mais précision 50 % et cas justes 60 % : 4 seuils de qualité sur 6
    r = _campagne_synthetique([5, 5, 5], fp=5, cas_justes=3)
    assert r["qualite"] == 4 and r["decision"] == "Apte"
    # Rappel 80 % à 100 % selon le passage : stabilité 20 points, plus précision et cas
    r = _campagne_synthetique([4, 5, 4], fp=5, cas_justes=3)
    assert r["indicateurs"]["I12"] == 20.0
    assert r["eliminatoires"] and r["qualite"] == 3 and r["decision"] == "Apte sous réserve"
    # Un éliminatoire manqué suffit : rappel 60 %
    assert _campagne_synthetique([3, 3, 3])["decision"] == "Non apte"
    # Passages échoués : I13 à 100 %
    assert _campagne_synthetique([5, 5, 5], echoue=True)["decision"] == "Non apte"


def test_commande_refuse_d_ecraser_un_classeur(tmp_path, capsys):
    chemin = tmp_path / "corrige.xlsx"
    assert ev.principal(["modele-corrige", str(chemin)]) == 0
    assert ev.principal(["modele-corrige", str(chemin)]) == 1
    assert "existe déjà" in capsys.readouterr().out
