"""Lot 3 — recette : RA-27 à RA-33, RA-59 à RA-63."""
import io
import struct
import zlib

import docx
from sqlalchemy import select

from qatag.db import SessionLocale
from qatag.models import EntreeJournal, Preuve

from .conftest import US_TEST, entetes, inscrire
from .test_cahier_cas import _valider_pv2
from .test_conception import _cas, _jusqua_pv2


def png(couleur=(200, 30, 30)) -> bytes:
    """Vraie image PNG 2×2 (lisible par python-docx)."""
    def bloc(nature, donnees):
        return (struct.pack(">I", len(donnees)) + nature + donnees +
                struct.pack(">I", zlib.crc32(nature + donnees) & 0xFFFFFFFF))
    lignes = b"".join(b"\x00" + bytes(couleur) * 2 for _ in range(2))
    return (b"\x89PNG\r\n\x1a\n" + bloc(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0)) +
            bloc(b"IDAT", zlib.compress(lignes)) + bloc(b"IEND", b""))


def _mission_validee(client, admin, ia):
    mid = _jusqua_pv2(client, admin, ia)
    _valider_pv2(client, admin, mid)
    return mid


def _executer(client, h, mid, cas_id, statut, obtenu="", commentaire="", captures=()):
    fichiers = [("captures", (nom, contenu, "image/png")) for nom, contenu in captures]
    return client.post("/api/missions/%d/cas/%d/executions" % (mid, cas_id), headers=h,
                       data={"statut": statut, "resultat_obtenu": obtenu,
                             "commentaire": commentaire}, files=fichiers or None)


def _recette(client, h, mid):
    r = client.get("/api/missions/%d/recette" % mid, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _statut(client, h, mid):
    return client.get("/api/missions/%d" % mid, headers=h).json()["statut"]


def _echec(client, h, mid, cas_id, obtenu="Formulaire de saisie affiché"):
    r = _executer(client, h, mid, cas_id, "FAIL", obtenu, captures=[("ecran.png", png())])
    assert r.status_code == 201, r.text
    return r


# --- RA-27 : ouverture -----------------------------------------------------------------

def test_recette_fermee_avant_pv2(client, admin, ia):
    mid = _jusqua_pv2(client, admin, ia)
    cas = _cas(client, admin, mid)
    r = _executer(client, admin, mid, cas["CT-001"]["id"], "PASS", "Message « Lien expiré »")
    assert r.status_code == 409 and r.json()["regle"] == "RA-27"


def test_premiere_execution_ouvre_la_recette(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    assert _statut(client, admin, mid) == "VALIDEE"
    cas = _cas(client, admin, mid)
    r = _executer(client, admin, mid, cas["CT-001"]["id"], "PASS", "Message « Lien expiré »")
    assert r.status_code == 201, r.text
    assert _statut(client, admin, mid) == "EN_RECETTE"
    etat = _recette(client, admin, mid)
    assert etat["compteurs"]["PASS"] == 1 and etat["compteurs"]["A_EXECUTER"] == 3


# --- RA-29, RA-59 : contenu d'une exécution ----------------------------------------------

def test_reussite_exige_le_resultat_obtenu(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    r = _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "PASS")
    assert r.status_code == 409 and r.json()["regle"] == "RA-29"


def test_echec_exige_une_capture(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    r = _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "FAIL",
                  "Formulaire affiché")
    assert r.status_code == 409 and r.json()["regle"] == "RA-59"


def test_non_testable_exige_un_commentaire(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cid = _cas(client, admin, mid)["CT-001"]["id"]
    r = _executer(client, admin, mid, cid, "BLOCKED")
    assert r.status_code == 409 and r.json()["regle"] == "RA-59"
    assert _executer(client, admin, mid, cid, "BLOCKED",
                     commentaire="Environnement de recette indisponible").status_code == 201


def test_statut_inconnu_refuse(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    r = _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "OK", "x")
    assert r.status_code == 409 and r.json()["regle"] == "RA-28"


# --- RA-31 : preuves ------------------------------------------------------------------------

def test_capture_doit_etre_une_vraie_image(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    r = _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "FAIL", "x",
                  captures=[("faux.png", b"<script>alert(1)</script>")])
    assert r.status_code == 409 and r.json()["regle"] == "RA-31"


def test_capture_conservee_en_base_avec_empreinte(client, admin, ia):
    import hashlib
    mid = _mission_validee(client, admin, ia)
    image = png()
    _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "FAIL", "x",
              captures=[("ecran.png", image)])
    preuve = _recette(client, admin, mid)["cas"][0]["executions"][0]["preuves"][0]
    assert preuve["empreinte"] == hashlib.sha256(image).hexdigest()
    r = client.get("/api/missions/%d/preuves/%d" % (mid, preuve["id"]), headers=admin)
    assert r.status_code == 200 and r.content == image
    assert r.headers["content-type"] == "image/png"


def test_capture_invisible_hors_de_la_mission(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    pid = _recette(client, admin, mid)["cas"][0]["executions"][0]["preuves"][0]["id"]
    inscrire(client, "autre@exemple.ci")
    r = client.get("/api/missions/%d/preuves/%d" % (mid, pid),
                   headers=entetes(client, "autre@exemple.ci"))
    assert r.status_code == 404


def test_historique_des_passages_conserve(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cid = _cas(client, admin, mid)["CT-001"]["id"]
    _echec(client, admin, mid, cid)
    _executer(client, admin, mid, cid, "PASS", "Message « Lien expiré »")
    ligne = _recette(client, admin, mid)["cas"][0]
    assert [e["statut"] for e in ligne["executions"]] == ["FAIL", "PASS"]
    assert ligne["statut"] == "PASS" and ligne["executions"][1]["passage"] == 2


# --- RA-30, RA-60 : anomalies et AG-09 ---------------------------------------------------------

def test_echec_ouvre_une_anomalie_qualifiee_par_ag09(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    a = _recette(client, admin, mid)["anomalies"][0]
    assert a["code"] == "ANO-001" and a["statut"] == "OUVERTE" and a["cas"] == "CT-001"
    assert a["gravite"] is None                          # rien n'est retenu sans l'humain
    assert a["agent"]["etat"] == "PROPOSEE" and a["agent"]["gravite"] == "Critique"
    assert a["agent"]["reproductibilite"] == "Non déterminée"
    assert "Formulaire de saisie affiché" in a["description"]
    consigne = ia.appels[-1]["corps"]["messages"][-1]["content"]
    assert "Message « Lien expiré »" in consigne and "Formulaire de saisie affiché" in consigne


def test_ag09_hors_format_echoue_proprement(client, admin, ia):
    ia.reponses["anomalie"] = "La gravité me semble élevée."
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    a = _recette(client, admin, mid)["anomalies"][0]
    assert a["agent"]["etat"] == "ECHEC" and "Anomalie" in a["agent"]["erreur"]
    assert _statut(client, admin, mid) == "EN_RECETTE"   # la recette continue


def test_ag09_systematique_refuse_sur_un_seul_echec(client, admin, ia):
    ia.reponses["anomalie"] = ia.reponses["anomalie"].replace("Non déterminée", "Systématique")
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    assert _recette(client, admin, mid)["anomalies"][0]["agent"]["etat"] == "ECHEC"


def _qualifier(client, h, mid, aid, gravite, motif=""):
    return client.put("/api/missions/%d/anomalies/%d/qualification" % (mid, aid), headers=h,
                      json={"gravite": gravite, "reproductibilite": "Non déterminée",
                            "motif": motif})


def test_confirmer_la_proposition_sans_motif(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    assert _qualifier(client, admin, mid, aid, "Critique").status_code == 200
    a = _recette(client, admin, mid)["anomalies"][0]
    assert a["gravite"] == "Critique" and a["qualifie_par"]["email"] == "bamk@exemple.ci"


def test_s_ecarter_de_la_proposition_exige_un_motif(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    r = _qualifier(client, admin, mid, aid, "Mineure")
    assert r.status_code == 409 and r.json()["regle"] == "RA-60"
    assert _qualifier(client, admin, mid, aid, "Mineure", "Libellé seulement").status_code == 200


def test_qualification_manuelle_sans_ia(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    with SessionLocale() as db:                  # l'IA disparaît après PV-2
        from qatag.models import ConfigIA
        db.query(ConfigIA).delete()
        db.commit()
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    a = _recette(client, admin, mid)["anomalies"][0]
    assert a["agent"]["etat"] is None
    assert _qualifier(client, admin, mid, a["id"], "Majeure").status_code == 200


def test_proposition_d_ag09_tracee_au_journal(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    with SessionLocale() as db:
        e = db.scalar(select(EntreeJournal).where(EntreeJournal.evenement == "anomalie_proposee"))
        assert e.agent == "AG-09" and '"Critique"' in e.details


# --- RA-61, RA-62 : cycle de vie ----------------------------------------------------------------

def _action(client, h, mid, aid, action, motif=""):
    return client.post("/api/missions/%d/anomalies/%d/statut" % (mid, aid), headers=h,
                       json={"action": action, "motif": motif})


def test_cycle_corrigee_puis_fermee_par_un_nouveau_passage_reussi(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cid = _cas(client, admin, mid)["CT-001"]["id"]
    _echec(client, admin, mid, cid)
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    assert _action(client, admin, mid, aid, "CORRIGER", "Livraison 1.2.1").status_code == 200
    assert _recette(client, admin, mid)["anomalies"][0]["statut"] == "CORRIGEE"
    _executer(client, admin, mid, cid, "PASS", "Message « Lien expiré »")
    a = _recette(client, admin, mid)["anomalies"][0]
    assert a["statut"] == "FERMEE" and "passage 2" in a["motif_statut"]


def test_echec_apres_correction_rouvre_sans_doublon(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cid = _cas(client, admin, mid)["CT-001"]["id"]
    _echec(client, admin, mid, cid)
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    _action(client, admin, mid, aid, "CORRIGER")
    _echec(client, admin, mid, cid, "Toujours le formulaire")
    anomalies = _recette(client, admin, mid)["anomalies"]
    assert len(anomalies) == 1 and anomalies[0]["statut"] == "OUVERTE"


def test_la_fermeture_ne_se_decide_pas(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    assert _action(client, admin, mid, aid, "FERMER").status_code == 409


def test_rejet_motive_puis_reouverture(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    assert _action(client, admin, mid, aid, "REJETER").status_code == 409
    assert _action(client, admin, mid, aid, "REJETER", "Erreur de saisie du testeur").status_code == 200
    assert _action(client, admin, mid, aid, "CORRIGER").status_code == 409
    assert _action(client, admin, mid, aid, "ROUVRIR", "Reproduit par le métier").status_code == 200


# --- RA-63 : clôture ----------------------------------------------------------------------------

def _cloturer(client, h, mid, decision, commentaire=""):
    return client.post("/api/missions/%d/recette/cloture" % mid, headers=h,
                       json={"decision": decision, "commentaire": commentaire})


def _tout_reussir(client, h, mid):
    for c in _cas(client, h, mid).values():
        assert _executer(client, h, mid, c["id"], "PASS", "Conforme").status_code == 201


def test_cloture_sans_reserve_quand_tout_est_reussi(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _tout_reussir(client, admin, mid)
    r = _cloturer(client, admin, mid, "CLOTURE")
    assert r.status_code == 200, r.text
    assert _statut(client, admin, mid) == "CLOTUREE"
    assert _recette(client, admin, mid)["cloture"]["auto_validation"] is True
    cid = _cas(client, admin, mid)["CT-001"]["id"]
    r = _executer(client, admin, mid, cid, "PASS", "x")
    assert r.status_code == 409 and r.json()["regle"] == "RA-63"


def test_cloture_sans_reserve_refusee_avec_la_liste(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    r = _cloturer(client, admin, mid, "CLOTURE")
    assert r.status_code == 409 and r.json()["regle"] == "RA-63"
    elements = r.json()["elements"]
    assert "CT-001 : échec" in elements and "CT-002 non exécuté" in elements
    assert any(e.startswith("ANO-001 ouverte, non qualifiée") for e in elements)


def test_cloture_avec_reserves_motivee_et_figee(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    assert _cloturer(client, admin, mid, "CLOTURE_AVEC_RESERVES").status_code == 409
    r = _cloturer(client, admin, mid, "CLOTURE_AVEC_RESERVES", "Mise en production acceptée par le métier")
    assert r.status_code == 200
    cloture = _recette(client, admin, mid)["cloture"]
    assert cloture["decision"] == "CLOTURE_AVEC_RESERVES" and len(cloture["reserves"]) == 5


def test_cloture_avec_reserves_refusee_sans_reserve(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _tout_reussir(client, admin, mid)
    r = _cloturer(client, admin, mid, "CLOTURE_AVEC_RESERVES", "x")
    assert r.status_code == 409


def test_cloture_reservee_au_superviseur_quand_plusieurs_participants(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    inscrire(client, "analyste@exemple.ci")
    client.post("/api/missions/%d/participants" % mid, json={"email": "analyste@exemple.ci"},
                headers=admin)
    h = entetes(client, "analyste@exemple.ci")
    _tout_reussir(client, h, mid)
    r = _cloturer(client, h, mid, "CLOTURE")
    assert r.status_code == 409 and r.json()["regle"] == "RA-12"
    assert "clôturer la recette" in r.json()["detail"]


# --- RA-32, RA-33 : cahier de recette --------------------------------------------------------------

def _cahier(client, h, mid):
    return client.get("/api/missions/%d/cahier-recette" % mid, headers=h)


def _texte(contenu: bytes) -> str:
    d = docx.Document(io.BytesIO(contenu))
    return "\n".join([p.text for p in d.paragraphs] +
                     [c.text for t in d.tables for r in t.rows for c in r.cells])


def test_cahier_refuse_avant_la_recette(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    r = _cahier(client, admin, mid)
    assert r.status_code == 409 and r.json()["regle"] == "RA-32"


def test_cahier_brouillon_depuis_la_base(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cas = _cas(client, admin, mid)
    _echec(client, admin, mid, cas["CT-001"]["id"])
    _executer(client, admin, mid, cas["CT-002"]["id"], "PASS", "Formulaire de saisie affiché")
    r = _cahier(client, admin, mid)
    assert r.status_code == 200, r.text
    assert "_brouillon" in r.headers["content-disposition"]
    texte = _texte(r.content)
    assert "CAHIER DE RECETTE" in texte and "DOCUMENT DE TRAVAIL" in texte
    assert "Synthèse de la recette" in texte and "ANO-001" in texte
    assert "☒  FAIL" in texte and "☒  PASS" in texte and "☒  NON EXÉCUTÉ" in texte
    d = docx.Document(io.BytesIO(r.content))
    assert len(d.inline_shapes) == 1                       # la capture de l'échec


def test_cahier_apres_cloture_et_journal(client, admin, ia):
    import hashlib
    mid = _mission_validee(client, admin, ia)
    _tout_reussir(client, admin, mid)
    _cloturer(client, admin, mid, "CLOTURE")
    r = _cahier(client, admin, mid)
    assert "_brouillon" not in r.headers["content-disposition"]
    assert "Recette clôturée sans réserve" in _texte(r.content)
    with SessionLocale() as db:
        e = db.scalar(select(EntreeJournal).where(
            EntreeJournal.evenement == "cahier_recette_exporte"))
        assert hashlib.sha256(r.content).hexdigest() in e.details
        assert '"conforme"' in e.details


def test_controle_ra33_bloque_un_document_incoherent(client, admin, ia, monkeypatch):
    from qatag import cahier_recette
    mid = _mission_validee(client, admin, ia)
    _executer(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"], "PASS", "x")
    monkeypatch.setitem(cahier_recette.CASES, "PASS", "FAIL")     # statut mal coché
    r = _cahier(client, admin, mid)
    assert r.status_code == 500 and r.json()["regle"] == "RA-33"
    with SessionLocale() as db:                                       # rien n'est tracé
        assert db.scalar(select(EntreeJournal).where(
            EntreeJournal.evenement == "cahier_recette_exporte")) is None


def test_preuve_jamais_dupliquee_sur_disque(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _echec(client, admin, mid, _cas(client, admin, mid)["CT-001"]["id"])
    with SessionLocale() as db:
        p = db.scalar(select(Preuve))
        assert p.chemin is None and p.taille > 0


def test_mission_us_test_disponible():
    assert "US-042" in US_TEST


# --- RA-64, RA-65 : suivi et levée des réserves -----------------------------------------------

def _clos_avec_reserves(client, h, mid):
    cas = _cas(client, h, mid)
    _echec(client, h, mid, cas["CT-001"]["id"])
    for code in ("CT-002", "CT-003", "CT-004"):
        _executer(client, h, mid, cas[code]["id"], "PASS", "Conforme")
    r = _cloturer(client, h, mid, "CLOTURE_AVEC_RESERVES", "Mise en production acceptée")
    assert r.status_code == 200, r.text
    return cas


def _lever(client, h, mid, commentaire=""):
    return client.post("/api/missions/%d/recette/levee-reserves" % mid, headers=h,
                       json={"commentaire": commentaire})


def test_seuls_les_cas_en_reserve_restent_executables(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cas = _clos_avec_reserves(client, admin, mid)
    r = _executer(client, admin, mid, cas["CT-002"]["id"], "PASS", "x")
    assert r.status_code == 409 and r.json()["regle"] == "RA-64"
    etat = _recette(client, admin, mid)
    assert etat["suivi_reserves"] is True
    assert {c["code"]: c["executable"] for c in etat["cas"]} == {
        "CT-001": True, "CT-002": False, "CT-003": False, "CT-004": False}


def test_levee_des_reserves_apres_correction_et_nouveau_passage(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cas = _clos_avec_reserves(client, admin, mid)
    r = _lever(client, admin, mid)
    assert r.status_code == 409 and r.json()["regle"] == "RA-65"
    assert any(e.startswith("ANO-001") for e in r.json()["elements"])
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    assert _action(client, admin, mid, aid, "CORRIGER", "Livraison 1.2.1").status_code == 200
    assert _executer(client, admin, mid, cas["CT-001"]["id"], "PASS",
                     "Message « Lien expiré »").status_code == 201
    etat = _recette(client, admin, mid)
    assert etat["anomalies"][0]["statut"] == "FERMEE" and etat["reserves_restantes"] == []
    assert _lever(client, admin, mid, "Corrections vérifiées en 1.2.1").status_code == 200
    assert _statut(client, admin, mid) == "RESERVES_LEVEES"
    cloture = _recette(client, admin, mid)["cloture"]
    assert cloture["decision"] == "CLOTURE_AVEC_RESERVES"              # décision inchangée
    assert len(cloture["reserves"]) == 2 and cloture["levee"]["auto_validation"] is True
    r = _executer(client, admin, mid, cas["CT-001"]["id"], "PASS", "x")
    assert r.status_code == 409 and r.json()["regle"] == "RA-63"
    texte = _texte(_cahier(client, admin, mid).content)
    assert "Réserves levées le" in texte and "Corrections vérifiées en 1.2.1" in texte
    with SessionLocale() as db:
        assert db.scalar(select(EntreeJournal).where(
            EntreeJournal.evenement == "reserves_levees")) is not None


def test_nouvel_echec_pendant_le_suivi_rouvre_l_anomalie(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cas = _clos_avec_reserves(client, admin, mid)
    aid = _recette(client, admin, mid)["anomalies"][0]["id"]
    _action(client, admin, mid, aid, "CORRIGER")
    _echec(client, admin, mid, cas["CT-001"]["id"], "Toujours accepté")
    assert _recette(client, admin, mid)["anomalies"][0]["statut"] == "OUVERTE"
    assert _lever(client, admin, mid).status_code == 409


def test_pas_de_levee_apres_une_cloture_sans_reserve(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    _tout_reussir(client, admin, mid)
    _cloturer(client, admin, mid, "CLOTURE")
    r = _lever(client, admin, mid)
    assert r.status_code == 409 and r.json()["regle"] == "RA-65"


def test_levee_reservee_au_superviseur_quand_plusieurs_participants(client, admin, ia):
    mid = _mission_validee(client, admin, ia)
    cas = _clos_avec_reserves(client, admin, mid)
    inscrire(client, "analyste@exemple.ci")
    client.post("/api/missions/%d/participants" % mid, json={"email": "analyste@exemple.ci"},
                headers=admin)
    h = entetes(client, "analyste@exemple.ci")
    aid = _recette(client, h, mid)["anomalies"][0]["id"]
    _action(client, h, mid, aid, "CORRIGER")
    _executer(client, h, mid, cas["CT-001"]["id"], "PASS", "Conforme")
    r = _lever(client, h, mid)
    assert r.status_code == 409 and r.json()["regle"] == "RA-12"
    assert "lever les réserves" in r.json()["detail"]
