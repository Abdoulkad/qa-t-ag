"""RA-46 : export du cahier des cas de test au gabarit Word DSI."""
import hashlib
import io

import docx

from .conftest import US_TEST
from .test_conception import _cas, _jusqua_pv2


def _texte(contenu: bytes) -> str:
    d = docx.Document(io.BytesIO(contenu))
    morceaux = [p.text for p in d.paragraphs]
    morceaux += [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join(morceaux)


def _exporter(client, h, mid):
    return client.get("/api/missions/%d/cahier-cas" % mid, headers=h)


def _valider_pv2(client, h, mid):
    cas = _cas(client, h, mid)
    for code in ("CT-003", "CT-004"):
        client.patch("/api/missions/%d/cas/%d" % (mid, cas[code]["id"]),
                     json={"donnees_test": "Abc1234 (7 caractères)",
                           "justification_technique": "Partition valide", "confirmer": True},
                     headers=h)
    r = client.post("/api/missions/%d/portes/PV-2" % mid, json={"decision": "APPROUVE"},
                    headers=h)
    assert r.status_code == 200, r.text


def test_refuse_tant_que_les_cas_ne_sont_pas_concus(client, admin, ia):
    r = client.post("/api/missions", json={"titre": "US-042", "texte": US_TEST}, headers=admin)
    r = _exporter(client, admin, r.json()["mission"]["id"])
    assert r.status_code == 409 and r.json()["regle"] == "RA-46"


def test_avant_pv2_document_de_travail_fidele_au_gabarit(client, admin, ia):
    mid = _jusqua_pv2(client, admin, ia)
    r = _exporter(client, admin, mid)
    assert r.status_code == 200
    assert "Cahier_cas_M-001_v1_brouillon.docx" in r.headers["content-disposition"]
    texte = _texte(r.content)
    assert "DOCUMENT DE TRAVAIL — cas non validés (PV-2 en attente)" in texte
    assert "{{" not in texte                                  # tous les jetons remplis
    for code in ("CT-001", "CT-002", "CT-003", "CT-004"):
        assert code + " : " in texte
    # recopie mot pour mot des cas enregistrés, sans reformulation
    assert "Lien émis à 08:00 / ouverture à 08:31" in texte
    assert "Message « Lien expiré »" in texte
    # CT-001 priorité Haute : seule la 2e case est cochée
    assert "☐  Critique        ☒  Haute        ☐  Moyenne" in texte
    assert "À_CLARIFIER" in texte                              # commentaire de CT-004
    d = docx.Document(io.BytesIO(r.content))
    blocs = [t for t in d.tables if t.rows and "Cas de test" in t.rows[0].cells[0].text]
    assert len(blocs) == 4                                     # un bloc par cas


def test_apres_pv2_valide_et_trace_au_journal(client, admin, ia):
    mid = _jusqua_pv2(client, admin, ia)
    _valider_pv2(client, admin, mid)
    r = _exporter(client, admin, mid)
    assert "Cahier_cas_M-001_v1.docx" in r.headers["content-disposition"]
    texte = _texte(r.content)
    assert "Cas validés à PV-2 le" in texte and "(auto-validation)" in texte
    assert "DOCUMENT DE TRAVAIL" not in texte and "À_CLARIFIER" not in texte
    entree = [e for e in client.get("/api/missions/%d/journal" % mid, headers=admin).json()
              if e["evenement"] == "cahier_cas_exporte"][-1]
    assert entree["details"]["empreinte"] == hashlib.sha256(r.content).hexdigest()
    assert entree["details"]["valide_pv2"] is True and entree["details"]["cas"] == 4


def test_priorite_hors_gabarit_refusee_avec_le_cas_en_cause(client, admin, ia):
    mid = _jusqua_pv2(client, admin, ia)
    cas = _cas(client, admin, mid)
    client.patch("/api/missions/%d/cas/%d" % (mid, cas["CT-002"]["id"]),
                 json={"priorite": "Basse"}, headers=admin)
    r = _exporter(client, admin, mid)
    assert r.status_code == 409 and r.json()["regle"] == "RA-46"
    assert r.json()["elements"] == ["CT-002 : priorité « Basse » absente du gabarit"]


def test_variantes_de_libelle_ramenees_au_gabarit():
    from qatag.scripts_qa import valeur_gabarit
    types = ("Fonctionnel", "Sécurité/Habilitation", "Interface utilisateur")
    assert valeur_gabarit("Sécurité / Habilitation", types) == "Sécurité/Habilitation"
    assert valeur_gabarit("haute", ("Critique", "Haute", "Moyenne")) == "Haute"
    assert valeur_gabarit("Performance", types) is None
