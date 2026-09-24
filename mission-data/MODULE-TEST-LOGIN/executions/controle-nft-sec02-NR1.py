#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Vérification NFT-SEC-02 : recherche des valeurs interdites (fichier .interdits)
dans TOUS les livrables QA du module. Rapporte uniquement le nombre d'occurrences
par fichier, jamais la valeur. Ne scanne pas le fichier .interdits lui-même."""
import os
from pathlib import Path

BASE = Path("mission-data/MODULE-TEST-LOGIN")
interdits_path = BASE / ".interdits"

valeurs = [l.strip("\r\n") for l in interdits_path.read_text(encoding="utf-8").splitlines() if l.strip()]

# extensions textuelles des livrables
EXT = {".json", ".jsonl", ".feature", ".md", ".txt", ".py", ".docx"}

trouve_total = 0
detail = []
for p in sorted(BASE.rglob("*")):
    if not p.is_file() or p.name == ".interdits":
        continue
    if p.suffix.lower() not in EXT:
        continue
    try:
        contenu = p.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        continue
    # .docx : extraction du texte
    if p.suffix.lower() == ".docx":
        try:
            import docx
            d = docx.Document(str(p))
            parts = [par.text for par in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    for c in row.cells:
                        parts.append(c.text)
            contenu = "\n".join(parts)
        except Exception:
            continue
    n = sum(contenu.count(v) for v in valeurs)
    if n > 0:
        trouve_total += n
        detail.append({"fichier": str(p.relative_to(BASE)), "occurrences": n})

print(f"nb_valeurs_interdites = {len(valeurs)}")
print(f"fichiers_avec_occurrence = {len(detail)}")
print(f"occurrences_totales = {trouve_total}")
for d in detail:
    print(f"  {d['fichier']} : {d['occurrences']}")
